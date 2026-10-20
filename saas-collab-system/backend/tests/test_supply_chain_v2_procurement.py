from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import RefreshToken

from apps.accounts.models import CustomUser
from apps.masterdata.models import SupplierMaster
from apps.permissions.models import DataScope
from apps.supply_chain import services
from apps.supply_chain.dates import business_date
from apps.supply_chain.archives import save_price
from apps.supply_chain.documents import IMPORT_HEADERS, IMPORT_SHEET, import_template, purchase_workbook, sku_labels
from apps.supply_chain.imports import commit_import, create_import, preview_excel
from apps.supply_chain.models import BusinessEvent, NotificationOutbox, PurchaseImport, PurchaseOrder, SupplierPriceHistory, SupplierProductPrice
from tests.test_supply_chain_v2_foundation import world, internal_user, supplier_user, po_payload, confirmed_order

BASE = "/api/internal/supply-chain/v2/"


@pytest.mark.django_db
def test_procurement_menu_migration_preserves_grants():
    from importlib import import_module
    from django.apps import apps
    from django.db import connection
    from apps.permissions.models import Permission, Role

    code = "menu.supply.supply_chain_purchase_orders.view"
    before = list(Role.permissions.through.objects.order_by("pk").values_list("pk", "role_id", "permission_id"))
    Permission.objects.filter(code=code).update(name="采购订单", metadata={"action_codes": ["purchasing.orders.view"]})
    import_module("apps.supply_chain.migrations.0004_procurement_menu").register(apps, connection.schema_editor())
    permission = Permission.objects.get(code=code)
    assert permission.permission_type == "menu" and permission.name == "采购协同"
    assert permission.metadata["action_codes"] == ["supply.v2.view"]
    assert before == list(Role.permissions.through.objects.order_by("pk").values_list("pk", "role_id", "permission_id"))


def test_business_date_accepts_beijing_today_before_utc_midnight(monkeypatch):
    from datetime import date, datetime, timezone as utc
    from rest_framework.exceptions import ValidationError as InputError
    from apps.supply_chain.serializers import PurchaseInput

    monkeypatch.setattr(timezone, "now", lambda: datetime(2026, 10, 12, 18, tzinfo=utc.utc))
    assert business_date() == date(2026, 10, 13)
    assert PurchaseInput().validate_order_date(date(2026, 10, 13)) == date(2026, 10, 13)
    with pytest.raises(InputError, match="下单日期不能晚于今天"):
        PurchaseInput().validate_order_date(date(2026, 10, 14))


def client(user):
    api = APIClient()
    api.force_authenticate(user=user)
    return api


def post(api, path, payload, key="request-key"):
    return api.post(BASE + path, payload, format="json", HTTP_IDEMPOTENCY_KEY=key)


def manual_payload(world, **kwargs):
    value = po_payload(world, **kwargs)
    value.pop("number")
    return value


def excel(rows):
    book = Workbook()
    sheet = book.active
    sheet.title = IMPORT_SHEET
    sheet.append(IMPORT_HEADERS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def row(world, group="A", price="2.5", ref="", qty=4):
    return [group, world["supplier"].code, world["country"].code, world["sku"].sku_code, qty, price, ref, str(business_date()), ""]


def test_real_api_create_edit_confirm_and_immutable_notification_snapshot(world):
    buyer = internal_user(world)
    api = client(buyer)
    payload = manual_payload(world, reference_number="R-SAME")
    created = post(api, "purchase-orders/", payload)
    assert created.status_code == 201 and created.data["success"]
    obj = created.data["data"]
    assert obj["number"].startswith("PO-") and obj["total_quantity"] == 10
    assert post(api, "purchase-orders/", payload).data["data"]["id"] == obj["id"]
    payload["lines"] = [{"sku_id": world["sku_b"].pk, "quantity": 5, "unit_price": "0"}]
    payload["expected_version"] = obj["version"]
    edited = api.patch(BASE + f"purchase-orders/{obj['id']}/", payload, format="json", HTTP_IDEMPOTENCY_KEY="edit")
    assert edited.status_code == 200
    assert edited.data["data"]["total_quantity"] == 5 and edited.data["data"]["total_amount"] == "0.0000"
    order = PurchaseOrder.objects.get(pk=obj["id"])
    assert order.lines.count() == 2 and order.lines.filter(is_active=True).count() == 1
    stale = post(api, f"purchase-orders/{obj['id']}/confirm/", {"expected_version": 1}, "stale")
    assert stale.status_code == 422 and not NotificationOutbox.objects.exists()
    placed = post(api, f"purchase-orders/{obj['id']}/confirm/", {"expected_version": 2}, "confirm")
    assert placed.status_code == 200 and placed.data["data"]["notification_status"] == "blocked"
    assert post(api, f"purchase-orders/{obj['id']}/confirm/", {"expected_version": 2}, "confirm").status_code == 200
    outbox = NotificationOutbox.objects.get()
    assert outbox.event.payload["purchase_snapshot"]["total_amount"] == "0.0000"
    metadata = api.patch(BASE + f"purchase-orders/{obj['id']}/metadata/", {"expected_version": 3, "reference_number": "CHANGED"}, format="json", HTTP_IDEMPOTENCY_KEY="meta")
    assert metadata.status_code == 200
    outbox.refresh_from_db()
    assert outbox.event.payload["purchase_snapshot"]["reference_number"] == "R-SAME"
    assert api.patch(BASE + f"purchase-orders/{obj['id']}/", {**payload, "expected_version": 4}, format="json", HTTP_IDEMPOTENCY_KEY="invalid-edit").status_code == 422


def test_duplicate_purchase_skus_are_rejected_with_all_row_numbers_on_create_and_update(world):
    buyer = internal_user(world)
    api = client(buyer)
    sku_id = world["sku"].pk
    duplicate_lines = [
        {"sku_id": sku_id, "quantity": 1, "unit_price": "2"},
        {"sku_id": world["sku_b"].pk, "quantity": 1, "unit_price": "3"},
        {"sku_id": sku_id, "quantity": 1, "unit_price": "2"},
    ]
    created = post(api, "purchase-orders/", manual_payload(world, lines=duplicate_lines), "duplicate-create")
    create_error = " ".join(created.data["data"]["detail"])
    assert created.status_code == 422 and all(number in create_error for number in ("1", "3"))
    assert not PurchaseOrder.objects.exists()

    order = post(api, "purchase-orders/", manual_payload(world, lines=[{"sku_id": sku_id, "quantity": 1}]), "valid-create").data["data"]
    duplicate_lines[0]["sku_id"] = world["sku_b"].pk
    duplicate_lines[2]["sku_id"] = world["sku_b"].pk
    updated = api.patch(BASE + f"purchase-orders/{order['id']}/", {
        **manual_payload(world, lines=duplicate_lines), "expected_version": order["version"]
    }, format="json", HTTP_IDEMPOTENCY_KEY="duplicate-update")
    update_error = " ".join(updated.data["data"]["detail"])
    assert updated.status_code == 422 and all(number in update_error for number in ("1", "3"))
    order_db = PurchaseOrder.objects.get(pk=order["id"])
    assert order_db.version == order["version"] and order_db.lines.filter(is_active=True).count() == 1


def test_default_prices_snapshot_manual_override_and_backdated_history(world):
    buyer = internal_user(world)
    sku = world["sku"]
    sku.purchase_price = Decimal("3")
    sku.save()
    api = client(buyer)
    first = post(api, "purchase-orders/", manual_payload(world, lines=[{"sku_id": sku.pk, "quantity": 2}]), "first").data["data"]
    assert Decimal(first["lines"][0]["unit_price"]) == Decimal("3")
    price = SupplierProductPrice.objects.get()
    maintained = post(api, "supplier-prices/", {"supplier_id": world["supplier"].pk, "sku_id": sku.pk,
        "unit_price": "4.25", "effective_date": str(business_date()), "expected_version": price.version}, "price")
    assert maintained.status_code == 200
    second = post(api, "purchase-orders/", manual_payload(world, lines=[{"sku_id": sku.pk, "quantity": 2}]), "second").data["data"]
    assert second["lines"][0]["product_reference_price"] == "3.0000"
    assert second["lines"][0]["supplier_reference_price"] == second["lines"][0]["unit_price"] == "4.2500"
    third = post(api, "purchase-orders/", manual_payload(world, order_date=str(business_date() - timedelta(days=1)), lines=[{"sku_id": sku.pk, "quantity": 2, "unit_price": "1"}]), "old").data["data"]
    assert third["lines"][0]["unit_price"] == "1.0000"
    price.refresh_from_db()
    assert price.unit_price == Decimal("4.25")
    assert SupplierPriceHistory.objects.filter(applied_to_current=False, unit_price=1).count() == 1
    # One manual edit replay never creates an extra history entry.
    before = SupplierPriceHistory.objects.count()
    assert post(api, "supplier-prices/", {"supplier_id": world["supplier"].pk, "sku_id": sku.pk,
        "unit_price": "4.25", "effective_date": str(business_date()), "expected_version": 1}, "price").status_code == 200
    assert SupplierPriceHistory.objects.count() == before


def test_blank_price_draft_cannot_notify_and_explicit_zero_is_valid(world):
    buyer = internal_user(world)
    api = client(buyer)
    created = post(api, "purchase-orders/", manual_payload(world, lines=[{"sku_id": world["sku"].pk, "quantity": 1, "unit_price": None}])).data["data"]
    assert created["total_amount"] is None and not created["prices_complete"]
    assert not SupplierProductPrice.objects.exists()
    assert post(api, f"purchase-orders/{created['id']}/confirm/", {"expected_version": 1}, "no-price").status_code == 422
    assert not NotificationOutbox.objects.exists()


@pytest.mark.parametrize("changes", [
    {"source_kind": "operating"}, {"source_kind": "import"}, {"supplier_id": 999999},
    {"lines": []}, {"lines": [{"sku_id": 999999, "quantity": 1}]},
    {"lines": [{"sku_id": 1, "quantity": 0}]},
    {"lines": [{"sku_id": 1, "quantity": 1, "unit_price": "-1"}]},
    {"lines": [{"sku_id": 1, "quantity": 1, "unit_price": "0.00001"}]},
])
def test_invalid_api_inputs_leave_no_partial_order_or_price(world, changes):
    buyer = internal_user(world)
    response = post(client(buyer), "purchase-orders/", {**manual_payload(world), **changes})
    assert response.status_code in {400, 422}
    assert not PurchaseOrder.objects.exists() and not SupplierProductPrice.objects.exists()


def test_scope_role_channel_and_supplier_file_isolation(world):
    buyer = internal_user(world)
    order = confirmed_order(world, buyer)
    own_supplier = supplier_user(world, world["supplier"], "own-supplier")
    other_supplier = supplier_user(world, world["supplier_b"], "other-supplier")
    for path in ["suppliers/", "supplier-prices/", "freight-companies/", "options/?kind=skus"]:
        assert client(own_supplier).get(BASE + path).status_code == 403
    assert client(own_supplier).get(BASE + f"purchase-orders/{order.pk}/").status_code == 403
    for suffix in ["excel/", "labels/"]:
        assert client(own_supplier).get(BASE + f"purchase-orders/{order.pk}/{suffix}").status_code == 403
        external = f"/api/external/supplier/supply-chain/v2/purchase-orders/{order.pk}/{suffix}"
        assert client(own_supplier).get(external).status_code == 200
        assert client(other_supplier).get(external).status_code == 404
        assert client(buyer).get(external).status_code == 403
    view_user = internal_user(world, "read-only", codes=("supply.v2.view",))
    assert post(client(view_user), "purchase-orders/", manual_payload(world), "readonly").status_code == 403
    custom = internal_user(world, "custom", scope=DataScope.ScopeType.CUSTOM, config={"supplier_ids": [world["supplier_b"].pk]})
    assert client(custom).get(BASE + f"purchase-orders/{order.pk}/").status_code == 404
    assert [x["id"] for x in client(custom).get(BASE + "suppliers/").data["data"]] == [world["supplier_b"].pk]
    token = RefreshToken.for_user(buyer).access_token
    token["channel"] = "miniapp"
    api = APIClient()
    api.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")
    assert api.get(BASE + "options/?kind=countries").status_code == 403


def test_supplier_and_freight_archives_versions_constraints_and_replay(world):
    buyer = internal_user(world)
    api = client(buyer)
    supplier_payload = {"code": "new-supplier", "name": "测试工厂", "province": "福建", "city": "厦门",
                        "factory_address": "工厂地址", "loading_address": "装柜地址", "tax_percent": "6.00", "notes": "备注"}
    created = post(api, "suppliers/", supplier_payload, "supplier-create")
    assert created.status_code == 201 and created.data["data"]["city"] == "厦门"
    assert post(api, "suppliers/", supplier_payload, "supplier-create").data["data"]["id"] == created.data["data"]["id"]
    payload = {**supplier_payload, "city": "泉州", "expected_version": 1}
    path = BASE + f"suppliers/{created.data['data']['id']}/"
    assert api.patch(path, payload, format="json", HTTP_IDEMPOTENCY_KEY="supplier-update").data["data"]["version"] == 2
    assert api.patch(path, payload, format="json", HTTP_IDEMPOTENCY_KEY="stale-supplier").status_code == 422
    assert api.patch(path, {**payload, "expected_version": 2, "tax_percent": 101}, format="json", HTTP_IDEMPOTENCY_KEY="invalid-tax").status_code == 400
    # Shared masterdata saves also advance the same version and reject stale instances.
    a = SupplierMaster.objects.get(pk=created.data["data"]["id"])
    b = SupplierMaster.objects.get(pk=a.pk)
    a.name = "新名称"
    a.save()
    with pytest.raises(ValidationError):
        b.save()
    freight = {"code": "forwarder", "name": "货运公司", "province": "福建", "city": "厦门", "address": "地址",
               "shipping_marks": "DF-US", "fee_per_cbm": "120.50", "tax_percent": "3.00", "notes": "备注"}
    carrier = post(api, "freight-companies/", freight, "freight-create")
    assert carrier.status_code == 201 and Decimal(carrier.data["data"]["fee_per_cbm"]) == Decimal("120.5")
    assert post(api, "freight-companies/", freight, "freight-create").data["data"]["id"] == carrier.data["data"]["id"]


def test_import_preview_commit_idempotency_refs_and_source_row_retention(world):
    buyer = internal_user(world)
    api = client(buyer)
    content = excel([row(world, "A", ref="REPEATED"), row(world, "B", ref="REPEATED")])
    response = api.post(BASE + "imports/", {"file": SimpleUploadedFile("orders.xlsx", content)}, HTTP_IDEMPOTENCY_KEY="upload")
    assert response.status_code == 201
    job = response.data["data"]
    assert job["preview"]["valid"] and job["preview"]["total_quantity"] == 8
    assert not PurchaseOrder.objects.exists()
    result = post(api, f"imports/{job['id']}/commit/", {"expected_version": 1}, "commit")
    assert result.status_code == 200 and len(result.data["data"]["order_ids"]) == 2
    assert post(api, f"imports/{job['id']}/commit/", {"expected_version": 1}, "commit").status_code == 200
    assert post(api, f"imports/{job['id']}/commit/", {"expected_version": 2}, "new-commit").status_code == 200
    assert PurchaseOrder.objects.count() == 2 and not NotificationOutbox.objects.exists()
    assert set(PurchaseOrder.objects.values_list("reference_number", flat=True)) == {"REPEATED"}
    assert set(PurchaseOrder.objects.values_list("status", flat=True)) == {"draft"}
    assert set(PurchaseOrder.objects.values_list("lines__import_row_number", flat=True)) == {2, 3}
    stored = PurchaseImport.objects.get()
    assert bytes(stored.file_content) == content
    other = internal_user(world, "other-buyer")
    assert client(other).get(BASE + f"imports/{stored.pk}/").status_code == 404


@pytest.mark.parametrize("mutation", ["unknown-sku", "zero-qty", "bad-price", "formula", "mixed-country", "conflicting-price"])
def test_import_errors_never_commit_partial_valid_rows(world, mutation):
    buyer = internal_user(world)
    rows = [row(world), row(world)]
    if mutation == "unknown-sku": rows[1][3] = "UNKNOWN"
    if mutation == "zero-qty": rows[1][4] = 0
    if mutation == "bad-price": rows[1][5] = "NaN"
    if mutation == "formula": rows[1][6] = "=1+1"
    if mutation == "mixed-country": rows[1][2] = world["country_ca"].code
    if mutation == "conflicting-price": rows[1][5] = "3"
    job = create_import(buyer, "bad.xlsx", excel(rows), idempotency_key="bad-upload")
    assert not job.preview["valid"] and job.preview["errors"][0]["row"] == 3
    with pytest.raises(ValidationError):
        commit_import(buyer, job.pk, expected_version=1, idempotency_key="bad-commit")
    assert not PurchaseOrder.objects.exists()


def test_import_rechecks_current_master_access_and_retains_preview_prices(world):
    buyer = internal_user(world)
    job = create_import(buyer, "orders.xlsx", excel([row(world, price="")]), idempotency_key="preview")
    assert job.preview["missing_prices_count"] == 1
    save_price(buyer, {"supplier_id": world["supplier"].pk, "sku_id": world["sku"].pk, "unit_price": Decimal(9),
                      "effective_date": business_date(), "expected_version": 0}, idempotency_key="price-change")
    committed = commit_import(buyer, job.pk, expected_version=1, idempotency_key="commit")
    assert PurchaseOrder.objects.get(pk=committed.order_ids[0]).lines.get().unit_price is None
    # A second preview cannot commit after its supplier is disabled.
    job2 = create_import(buyer, "orders.xlsx", excel([row(world)]), idempotency_key="preview2")
    supplier = SupplierMaster.objects.get(pk=world["supplier"].pk)
    supplier.status = "inactive"
    supplier.save()
    with pytest.raises(ValidationError):
        commit_import(buyer, job2.pk, expected_version=1, idempotency_key="commit2")
    assert PurchaseOrder.objects.count() == 1


def test_download_files_typed_values_formula_injection_and_labels(world):
    fitz = pytest.importorskip("fitz")
    buyer = internal_user(world)
    order = confirmed_order(world, buyer, payload=po_payload(world, reference_number="=HYPERLINK(\"malicious\")"))
    snapshot = services.purchase_snapshot(order)
    workbook = load_workbook(BytesIO(purchase_workbook(snapshot)), data_only=False)
    assert workbook.sheetnames == ["采购单", "发货数量填写", "箱号信息填写"]
    assert workbook["采购单"]["B3"].data_type == "s" and workbook["采购单"]["B3"].value.startswith("=")
    assert workbook["采购单"]["D8"].value == 10 and workbook["采购单"]["E8"].value == 2.5
    assert workbook["采购单"]["F8"].data_type == "f"
    pdf = fitz.open(stream=sku_labels(snapshot), filetype="pdf")
    assert len(pdf) == 1 and world["sku"].sku_code in pdf[0].get_text()
    assert abs(pdf[0].rect.width - 60 * 72 / 25.4) < 0.1
    # An untouched downloadable template is syntactically readable with zero data rows.
    preview = preview_excel(buyer, import_template())
    assert preview["rows_count"] == 0 and not preview["valid"]
    api = client(buyer)
    for suffix in ["excel/", "labels/"]:
        response = api.get(BASE + f"purchase-orders/{order.pk}/{suffix}")
        assert response.status_code == 200 and response.streaming
        assert "attachment" in response["Content-Disposition"]


def test_pagination_search_cancel_and_required_idempotency(world):
    buyer = internal_user(world)
    api = client(buyer)
    assert api.post(BASE + "purchase-orders/", manual_payload(world), format="json").status_code == 400
    a = post(api, "purchase-orders/", manual_payload(world, reference_number="FIND-A"), "A").data["data"]
    post(api, "purchase-orders/", manual_payload(world, reference_number="FIND-B"), "B")
    result = api.get(BASE + "purchase-orders/?search=FIND&limit=1&offset=1")
    assert result.data["count"] == 2 and len(result.data["data"]) == 1
    assert api.get(BASE + "purchase-orders/?limit=1000").status_code == 400
    assert post(api, f"purchase-orders/{a['id']}/cancel/", {"expected_version": 1}, "cancel").status_code == 200
    assert post(api, f"purchase-orders/{a['id']}/confirm/", {"expected_version": 2}, "after-cancel").status_code == 422
    assert not NotificationOutbox.objects.exists()


def test_price_history_append_advances_version_and_shared_manual_edits_need_supplier_scope(world):
    buyer = internal_user(world)
    order = confirmed_order(world, buyer)
    price = SupplierProductPrice.objects.get()
    version = price.version
    api = client(buyer)
    payload = {"supplier_id": price.supplier_id, "sku_id": price.sku_id, "unit_price": "1",
               "effective_date": str(price.effective_date - timedelta(days=2)), "expected_version": version}
    response = post(api, "supplier-prices/", payload, "historical-price")
    assert response.status_code == 200 and response.data["data"]["version"] == version + 1
    price.refresh_from_db()
    assert price.unit_price == Decimal("2.5")
    payload["effective_date"] = str(business_date())
    assert post(api, "supplier-prices/", payload, "stale-price").status_code == 422
    own = internal_user(world, "own-archive", scope=DataScope.ScopeType.OWN)
    assert post(client(own), "supplier-prices/", {**payload, "expected_version": price.version}, "own-price").status_code == 403
    assert post(client(own), "suppliers/", {"code": "own-new", "name": "Shared"}, "own-supplier").status_code == 403
    custom = internal_user(world, "assigned-archive", scope=DataScope.ScopeType.CUSTOM, config={"supplier_ids": [price.supplier_id]})
    assert post(client(custom), "supplier-prices/", {**payload, "expected_version": price.version}, "assigned-price").status_code == 200


def test_imported_draft_edit_keeps_source_and_original_rows(world):
    buyer = internal_user(world)
    job = create_import(buyer, "orders.xlsx", excel([row(world)]), idempotency_key="edit-import")
    job = commit_import(buyer, job.pk, expected_version=1, idempotency_key="edit-import-commit")
    order = PurchaseOrder.objects.get(pk=job.order_ids[0])
    payload = {"supplier_id": order.supplier_id, "country_id": order.country_id, "source_kind": "import", "expected_version": 1,
        "reference_number": "NEW-REF", "lines": [
            {"sku_id": world["sku"].pk, "quantity": 3, "unit_price": "3", "import_row_number": 2},
            {"sku_id": world["sku_b"].pk, "quantity": 1, "unit_price": "4", "import_row_number": None},
        ]}
    response = client(buyer).patch(BASE + f"purchase-orders/{order.pk}/", payload, format="json", HTTP_IDEMPOTENCY_KEY="edit-import-draft")
    assert response.status_code == 200 and response.data["data"]["source_kind"] == "import", response.data
    assert [(line["sku_id"], line["import_row_number"]) for line in response.data["data"]["lines"]] == [
        (world["sku"].pk, 2), (world["sku_b"].pk, None)]
    assert order.lines.count() == 3

    # A second edit can round-trip both imported and hand-added rows.
    payload["expected_version"] = 2
    payload["lines"][1]["quantity"] = 2
    response = client(buyer).patch(BASE + f"purchase-orders/{order.pk}/", payload, format="json", HTTP_IDEMPOTENCY_KEY="edit-import-draft-again")
    assert response.status_code == 200
    assert [(line["sku_id"], line["import_row_number"]) for line in response.data["data"]["lines"]] == [
        (world["sku"].pk, 2), (world["sku_b"].pk, None)]

    # Existing imported SKUs cannot be assigned invented row numbers.
    forged = {**payload, "expected_version": 3, "lines": [
        {"sku_id": world["sku"].pk, "quantity": 3, "unit_price": "3", "import_row_number": 2},
        {"sku_id": world["sku_b"].pk, "quantity": 1, "unit_price": "4", "import_row_number": 3},
    ]}
    assert client(buyer).patch(BASE + f"purchase-orders/{order.pk}/", forged, format="json", HTTP_IDEMPOTENCY_KEY="edit-import-forged").status_code == 422
    duplicate = {**payload, "expected_version": 3, "lines": payload["lines"] + [payload["lines"][1]]}
    assert client(buyer).patch(BASE + f"purchase-orders/{order.pk}/", duplicate, format="json", HTTP_IDEMPOTENCY_KEY="edit-import-duplicate").status_code == 422

    dropped = {**payload, "expected_version": 3, "lines": payload["lines"][1:]}
    response = client(buyer).patch(BASE + f"purchase-orders/{order.pk}/", dropped, format="json", HTTP_IDEMPOTENCY_KEY="edit-import-drop")
    assert response.status_code == 200 and [(line["sku_id"], line["import_row_number"]) for line in response.data["data"]["lines"]] == [
        (world["sku_b"].pk, None)]


def test_invalid_xlsx_structure_and_resource_limits_fail_cleanly(world):
    from zipfile import ZipFile
    buyer = internal_user(world)
    with pytest.raises(ValidationError):
        preview_excel(buyer, b"not-xlsx")
    with pytest.raises(ValidationError):
        preview_excel(buyer, b"x" * (5 * 1024 * 1024 + 1))
    raw = BytesIO()
    with ZipFile(raw, "w") as archive:
        archive.writestr("xl/workbook.xml", "<!DOCTYPE test [<!ENTITY x 'xx'>]><bad>&x;</bad>")
        archive.writestr("[Content_Types].xml", "<bad")
    with pytest.raises(ValidationError):
        preview_excel(buyer, raw.getvalue())
