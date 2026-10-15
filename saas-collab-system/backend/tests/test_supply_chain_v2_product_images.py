from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image

from apps.supply_chain.documents import purchase_workbook
from apps.supply_chain import supplier_product_images as product_images


def png_bytes(size=(24, 18)):
    output = BytesIO()
    Image.new("RGB", size, "#2e8b57").save(output, format="PNG")
    return output.getvalue()


def snapshot(image_url="/media/product-images/tenant-7/item.png"):
    return {
        "number": "PO-001", "supplier_name": "供方", "country_name": "美国",
        "reference_number": "R-1", "order_date": "2026-10-14", "total_quantity": 2,
        "total_amount": "12.0000", "notes": "", "lines": [{
            "sku_id": 41, "sku": "SKU-1", "name": "商品", "quantity": 2,
            "unit_price": "6.0000", "product_details": {
                "image_url": image_url, "spu": "=1+1", "brand": "品牌",
                "category": "分类", "specification": "规格", "spec_values": ["A", "B"],
                "color": "红", "size": "M", "material": "棉", "unit": "件",
                "package_length_cm": 10, "package_width_cm": 11, "package_height_cm": 12,
                "package_weight_g": 25, "package_volume_cbm": "0.001",
                "origin_country": "CN", "hs_code": "123456", "inventory_type": "普通",
                "product_description": "描述", "selling_points": ["耐用", "轻便"], "product_type": "成品",
            },
        }],
    }


def test_purchase_workbook_embeds_thumbnail_and_exports_text_metadata():
    workbook_snapshot = snapshot("https://assets.example/item.png")
    content = purchase_workbook(workbook_snapshot, images={41: {"body": png_bytes(), "status": "available"}})
    with ZipFile(BytesIO(content)) as archive:
        media = [name for name in archive.namelist() if name.startswith("xl/media/")]
        assert len(media) == 1
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["采购单", "发货数量填写", "箱号信息填写"]
    purchase = workbook["采购单"]
    assert [purchase.cell(8, col).value for col in range(1, 6)] == ["SKU-1", "商品", 2, 6, '=IF(D8="","",C8*D8)']
    assert len(purchase._images) == 1 and purchase._images[0].anchor._from.row == 7
    assert purchase.row_dimensions[8].height == 85
    assert purchase["H8"].value == "=1+1" and purchase["H8"].data_type == "s"
    assert purchase["L8"].value == '["A", "B"]' and purchase["L8"].data_type == "s"
    assert purchase["Q8"].value == "10" and purchase["Q8"].data_type == "s"
    assert purchase["G8"].hyperlink.target == workbook_snapshot["lines"][0]["product_details"]["image_url"]


def test_missing_and_unavailable_images_are_visible_without_unsafe_links():
    snap = snapshot("javascript:alert(1)")
    snap["lines"].append({**snap["lines"][0], "sku_id": 42, "product_details": {}})
    content = purchase_workbook(snap, images={41: {"body": None, "status": "failed"}, 42: {"body": None, "status": "missing"}})
    purchase = load_workbook(BytesIO(content))["采购单"]
    assert purchase["F8"].value == "图片读取失败"
    assert purchase["G8"].value is None and purchase["G8"].hyperlink is None
    assert purchase["F9"].value == "暂无图片"


def test_loader_reencodes_local_image_and_rejects_cross_tenant_and_unsafe_remote(monkeypatch):
    class Storage:
        def exists(self, path):
            return path == "product-images/tenant-7/item.png"

        def open(self, path, mode):
            assert path == "product-images/tenant-7/item.png"
            return BytesIO(png_bytes((1000, 500)))

    storage = Storage()
    monkeypatch.setattr(product_images, "default_storage", storage)
    monkeypatch.setattr("apps.products.views.default_storage", storage)
    good = product_images.load_product_images(snapshot(), tenant_id=7)[41]
    assert good["status"] == "available"
    with Image.open(BytesIO(good["body"])) as image:
        assert image.format == "JPEG" and image.size == (320, 160)
    wrong_tenant = product_images.load_product_images(snapshot(), tenant_id=8)[41]
    assert wrong_tenant == {"body": None, "status": "failed"}
    unsafe = product_images.load_product_images(snapshot("data:image/png;base64,AAAA"), tenant_id=7)[41]
    assert unsafe == {"body": None, "status": "failed"}
