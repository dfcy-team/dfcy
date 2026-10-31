from io import BytesIO

import pytest
from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from apps.supply_chain import documents
from apps.supply_chain.documents import purchase_workbook
from apps.supply_chain.supplier_products import enrich_snapshot
from apps.supply_chain import costs, services
from tests.test_supply_chain_v2_foundation import (confirmed_order, internal_user, supplier_user,
    po_payload, shipment_payload, world)
from tests.test_supply_chain_v2_shipments import EXTERNAL, INTERNAL, supplier_client, buyer_client


def test_legacy_sku_enrichment_uses_related_sku_and_preserves_frozen_blank(world):
    world['sku'].legacy_sku_code = 'OLD-42'
    world['sku'].save(update_fields=['legacy_sku_code'])
    original = {'lines': [
        {'id': 1, 'sku_id': world['sku'].pk, 'product_details': {'image_url': ''}},
        {'id': 2, 'sku_id': world['sku'].pk, 'legacy_sku_code': '', 'product_details': {'image_url': ''}},
    ]}
    enriched = enrich_snapshot(original, world['tenant'].pk)
    assert [line['legacy_sku_code'] for line in enriched['lines']] == ['OLD-42', '']


def test_purchase_workbook_places_legacy_sku_after_current_sku_and_preserves_fields():
    content = purchase_workbook({
        'number': 'PO-1', 'supplier_name': '供方', 'country_name': '美国', 'reference_number': '',
        'order_date': '2026-10-14', 'total_quantity': 3, 'total_amount': None, 'notes': '',
        'lines': [{'sku_id': 1, 'sku': 'SKU-1', 'legacy_sku_code': 'OLD-1', 'name': '商品',
                   'quantity': 3, 'unit_price': None, 'product_details': {}}],
    })
    workbook = load_workbook(BytesIO(content))
    purchase = workbook['采购单']
    quantity = workbook['发货数量填写']
    assert [purchase.cell(7, col).value for col in range(1, 9)] == [
        '商品SKU', '旧SKU', '商品名称', '采购数量', '采购单价', '金额', '商品图片', '图片地址']
    assert [purchase.cell(8, col).value for col in range(1, 7)] == ['SKU-1', 'OLD-1', '商品', 3, None, '=IF(E8="","",D8*E8)']
    assert [quantity.cell(5, col).value for col in range(1, 7)] == [
        '商品SKU', '旧SKU', '商品名称', '采购数量', '本批发货数量', '备注']
    assert [quantity.cell(6, col).value for col in range(1, 7)] == ['SKU-1', 'OLD-1', '商品', 3, None, None]
    assert quantity['B5'].fill.fgColor.rgb == '0028465C'
    assert quantity['E6'].fill.fgColor.rgb != quantity['D6'].fill.fgColor.rgb


def test_inbound_reconciliation_workbook_headers_values_null_zero_and_split_lines():
    from apps.supply_chain.documents import inbound_reconciliation_workbook
    workbook = load_workbook(BytesIO(inbound_reconciliation_workbook({
        'number': 'IN-1', 'batch_number': 'BATCH-1', 'lines': [
            {'sku_id': 1, 'legacy_sku_code': 'OLD-1', 'sku': 'NEW-1', 'product_name': '商品', 'expected_quantity': 5, 'received_quantity': 0},
            {'sku_id': 1, 'legacy_sku_code': '', 'sku': 'NEW-1', 'product_name': '商品', 'expected_quantity': 2, 'received_quantity': None},
        ]
    }, images={1: {'status': 'missing'}})))
    sheet = workbook.active
    assert [c.value for c in sheet[1]] == ['序号', '清点单号', '采购单号', '商品图片', '商品名称', 'SKU', '发货数量', '实际清点数量', '差异数量']
    assert [c.value for c in sheet[2]] == [1, 'IN-1', 'BATCH-1', '暂无图片', '商品', 'OLD-1', 5, None, None]
    assert [c.value for c in sheet[3]] == [2, 'IN-1', 'BATCH-1', '暂无图片', '商品', 'NEW-1', 2, None, None]


def test_sku_labels_use_trimmed_legacy_code_for_qr_and_printed_code_then_fallback(monkeypatch):
    fitz = __import__('pytest').importorskip('fitz')
    payloads = []
    actual_widget = documents.QrCodeWidget

    def recording_widget(payload, **kwargs):
        payloads.append(payload)
        return actual_widget(payload, **kwargs)

    monkeypatch.setattr(documents, 'QrCodeWidget', recording_widget)
    pdf = fitz.open(stream=documents.sku_labels({'lines': [
        {'sku': 'NEW-1', 'legacy_sku_code': '  OLD-1  ', 'name': '商品一'},
        {'sku': 'NEW-2', 'legacy_sku_code': '   ', 'name': '商品二'},
    ]}), filetype='pdf')

    assert payloads == ['OLD-1', 'NEW-2']
    assert len(pdf) == 2
    assert 'OLD-1' in pdf[0].get_text() and '商品一' in pdf[0].get_text()
    assert 'NEW-2' in pdf[1].get_text() and '商品二' in pdf[1].get_text()


def test_shipment_image_route_allows_document_extra_inactive_sku_and_isolates_supplier(world, monkeypatch):
    from apps.supply_chain import internal_product_images

    buyer = internal_user(world, 'image-buyer')
    order = confirmed_order(world, buyer)
    world['sku_b'].image_url = f"/media/product-images/tenant-{world['tenant'].pk}/extra.png"
    world['sku_b'].save(update_fields=['image_url'])
    shipment = services.create_supplier_shipment(buyer, shipment_payload(world, order, 'image-extra',
        lines=[{'sku_id': world['sku_b'].pk, 'quantity': 1}]), idempotency_key='image-extra')
    monkeypatch.setattr(internal_product_images, 'load_product_images',
        lambda snapshot, tenant_id: {world['sku_b'].pk: {'body': b'image', 'status': 'available'}})
    path = f"{INTERNAL}supplier-shipments/{shipment.pk}/skus/{world['sku_b'].pk}/image/"
    from apps.supply_chain.shipments import snapshot as shipment_snapshot
    line = shipment_snapshot(shipment, internal=True)['lines'][0]
    assert line['product_details']['image_url'] == world['sku_b'].image_url
    assert line['image_preview_url'] == path
    assert buyer_client(buyer).get(path).status_code == 200
    supplier = supplier_user(world, world['supplier'], 'image-supplier')
    option = supplier_client(supplier).get(f"{EXTERNAL}shipment-options/").data['data']['skus']
    sku_option = next(item for item in option if item['id'] == world['sku_b'].pk)
    assert sku_option['legacy_sku_code'] == '' and sku_option['image_preview_url'].endswith(
        f"/shipment-options/skus/{world['sku_b'].pk}/image/")
    assert supplier_client(supplier).get(sku_option['image_preview_url']).status_code == 200
    supplier_path = f"{EXTERNAL}shipments/{shipment.pk}/skus/{world['sku_b'].pk}/image/"
    response = supplier_client(supplier).get(supplier_path)
    assert response.status_code == 200 and response['Cache-Control'] == 'private, no-store'
    assert supplier_client(supplier_user(world, world['supplier_b'], 'other-image-supplier')).get(supplier_path).status_code == 404
    world['sku_b'].is_active = False
    world['sku_b'].save(update_fields=['is_active'])
    assert buyer_client(buyer).get(path).status_code == 200


def test_display_only_sku_edits_do_not_change_cost_input_fingerprint(world):
    buyer = internal_user(world, 'fingerprint-buyer')
    order = confirmed_order(world, buyer)
    shipment = services.create_supplier_shipment(buyer, shipment_payload(world, order, 'fingerprint'),
        idempotency_key='fingerprint')
    before = costs.inputs('shipment', shipment)[1]
    world['sku'].legacy_sku_code = 'OLD-FINGERPRINT'
    world['sku'].image_url = '/media/product-images/tenant-display-only/image.png'
    world['sku'].save(update_fields=['legacy_sku_code', 'image_url'])
    after = costs.inputs('shipment', shipment)[1]
    assert before == after


@pytest.mark.django_db
@pytest.mark.parametrize("readded_row_number", [None, 99])
def test_imported_sku_row_provenance_survives_removal_and_readdition(world, readded_row_number):
    buyer = internal_user(world, f"import-provenance-{readded_row_number}")
    order = services.create_purchase_order(buyer, po_payload(world, source="import", number="PO-PROVENANCE",
        import_batch_ref="batch-provenance", lines=[{"sku_id": world["sku"].pk, "quantity": 2,
            "unit_price": "1", "import_row_number": 8}]), idempotency_key="provenance-create")
    edit = {"supplier_id": world["supplier"].pk, "country_id": world["country"].pk,
        "source_kind": "import", "lines": [{"sku_id": world["sku_b"].pk, "quantity": 1,
            "unit_price": "1", "import_row_number": None}]}

    order = services.update_purchase_order(buyer, order.pk, edit, expected_version=1,
        idempotency_key="provenance-remove-original")
    assert order.version == 2

    readded = {**edit, "lines": [{"sku_id": world["sku"].pk, "quantity": 2,
        "unit_price": "1", "import_row_number": readded_row_number}]}
    with pytest.raises(ValidationError):
        services.update_purchase_order(buyer, order.pk, readded, expected_version=2,
            idempotency_key=f"provenance-readd-invalid-{readded_row_number}")

    correct = {**edit, "lines": [{"sku_id": world["sku"].pk, "quantity": 2,
        "unit_price": "1", "import_row_number": 8}]}
    order = services.update_purchase_order(buyer, order.pk, correct, expected_version=2,
        idempotency_key=f"provenance-readd-correct-{readded_row_number}")
    assert order.version == 3
    assert list(order.lines.filter(is_active=True).values_list("sku_id", "import_row_number")) == [(world["sku"].pk, 8)]

    # A genuinely new SKU remains unattributed even across subsequent edits.
    new_sku = {**edit, "lines": [{"sku_id": world["sku_b"].pk, "quantity": 1,
        "unit_price": "1", "import_row_number": None}]}
    order = services.update_purchase_order(buyer, order.pk, new_sku, expected_version=3,
        idempotency_key=f"provenance-new-sku-{readded_row_number}")
    order = services.update_purchase_order(buyer, order.pk, new_sku, expected_version=4,
        idempotency_key=f"provenance-new-sku-again-{readded_row_number}")
    assert list(order.lines.filter(is_active=True).values_list("sku_id", "import_row_number")) == [(world["sku_b"].pk, None)]


def test_batch_and_inbound_image_routes_enforce_document_and_actor_scope(world, monkeypatch):
    from apps.supply_chain import internal_product_images
    from tests.test_supply_chain_v2_logistics import BASE, arrive, client, create, generate, setup

    world['sku'].image_url = f"/media/product-images/tenant-{world['tenant'].pk}/document.png"
    world['sku'].save(update_fields=['image_url'])
    buyer, shipment, _other, carrier, warehouse = setup(world)
    monkeypatch.setattr(internal_product_images, 'load_product_images',
        lambda snapshot, tenant_id: {world['sku'].pk: {'body': b'image', 'status': 'available'}})
    api = client(buyer)
    batch, _ = create(api, [shipment], carrier, 'image-batch')
    batch_path = f"{BASE}dispatch-batches/{batch['id']}/skus/{world['sku'].pk}/image/"
    response = api.get(batch_path)
    assert response.status_code == 200 and response['Cache-Control'] == 'private, no-store'
    assert api.get(f"{BASE}dispatch-batches/{batch['id']}/skus/{world['sku_b'].pk}/image/").status_code == 404

    batch = arrive(api, batch)
    inbound = generate(api, batch, warehouse)
    inbound_path = f"{BASE}inbound-orders/{inbound['id']}/skus/{world['sku'].pk}/image/"
    response = api.get(inbound_path)
    assert response.status_code == 200 and response['Cache-Control'] == 'private, no-store'
    assert api.get(f"{BASE}inbound-orders/{inbound['id']}/skus/{world['sku_b'].pk}/image/").status_code == 404

    denied = client(internal_user(world, 'image-scope-denied', scope='own'))
    assert denied.get(batch_path).status_code == 404
    assert denied.get(inbound_path).status_code == 404


def test_inbound_export_embeds_image_and_keeps_formula_like_text_literal():
    from PIL import Image
    from apps.supply_chain.documents import inbound_reconciliation_workbook
    stream = BytesIO()
    Image.new('RGB', (20, 10), 'white').save(stream, format='PNG')
    payload = {'number': 'IN-2', 'batch_number': 'B-2', 'lines': [
        {'sku_id': 1, 'sku': '=BAD()', 'product_name': '=NAME()', 'expected_quantity': 2, 'received_quantity': 5}]}
    book = load_workbook(BytesIO(inbound_reconciliation_workbook(payload, images={1: {'body': stream.getvalue()}})))
    sheet = book.active
    assert len(sheet._images) == 1
    assert sheet['I2'].value is None
    assert sheet['F2'].data_type == 's' and sheet['E2'].data_type == 's'
    assert sheet['E2'].alignment.wrap_text
