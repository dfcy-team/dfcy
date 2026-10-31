from io import BytesIO
from pathlib import Path
import pytest
from openpyxl import load_workbook
from apps.supply_chain.bigseller import export_batch
from apps.supply_chain.documents import inbound_reconciliation_workbook
from apps.supply_chain.models import DispatchBatch
from tests.test_supply_chain_v2_foundation import world, supplier_user
from tests.test_supply_chain_v2_logistics import setup, create, client, BASE, arrive, generate, checked, write

@pytest.mark.django_db
def test_bigseller_export_preserves_template_and_batch_quantities(world):
    buyer, first, second, carrier, _ = setup(world)
    batch, _ = create(client(buyer), [first, second], carrier)
    obj=DispatchBatch.objects.get(pk=batch['id'])
    before=load_workbook(Path(__file__).parents[1]/'apps/supply_chain/templates/bigseller-purchase-zh.xlsx')
    workbook=load_workbook(BytesIO(export_batch(obj)))
    assert workbook.sheetnames == before.sheetnames
    assert [c.value for c in workbook['SKU'][1]] == [c.value for c in before['SKU'][1]]
    sheet=workbook['SKU']
    assert sheet['A2'].value == obj.number
    assert sheet['C2'].value == sum(l['quantity'] for sh in batch['shipments'] for l in sh['lines'])
    assert sheet['I2'].value == '国内供应商'
    assert sheet['F2'].value == 'CNY' and sheet['G2'].value == 1
    response=client(buyer).get(BASE+f'dispatch-batches/{obj.pk}/bigseller/')
    assert response.status_code == 400

    external=supplier_user(world,world['supplier'],'bigseller-external')
    assert client(external).get(BASE+f'dispatch-batches/{obj.pk}/bigseller/').status_code == 403

@pytest.mark.django_db
def test_bigseller_partial_allocation_and_cancelled_guard(world):
    from apps.supply_chain import logistics
    buyer, first, second, carrier, _=setup(world)
    api=client(buyer)
    lines=list(first.lines.all())
    body={'shipment_ids':[first.pk],'shipment_versions':[{'id':first.pk,'version':first.version}],
          'freight_company_id':carrier.pk,'freight_number':'TEST',
          'allocations':[{'shipment_line_id':lines[0].pk,'quantity':1}]}
    response=api.post(BASE+'dispatch-batches/',body,format='json',HTTP_IDEMPOTENCY_KEY='partial-export')
    assert response.status_code==201,response.data
    obj=DispatchBatch.objects.get(pk=response.data['data']['id'])
    sheet=load_workbook(BytesIO(export_batch(obj)))['SKU']
    assert sheet['C2'].value == 1
    logistics.cancel_batch(buyer,obj.pk,{'expected_version':obj.version,'reason':'test'},'cancel-export')
    assert api.get(BASE+f'dispatch-batches/{obj.pk}/bigseller/').status_code==400


@pytest.mark.django_db
def test_bigseller_sku_mode_weighted_price_and_query_bound(world, django_assert_num_queries):
    from apps.supply_chain.models import ShipmentLine, domain_write
    buyer, first, second, carrier, _ = setup(world)
    first_line=first.lines.first();second_line=second.lines.first()
    sku=first_line.sku;sku.legacy_sku_code='BS-OLD';sku.save(update_fields=['legacy_sku_code'])
    with domain_write():
        first_line.unit_price=2;first_line.save(update_fields=['unit_price'])
        second_line.unit_price=4;second_line.save(update_fields=['unit_price'])
    batch,_=create(client(buyer),[first,second],carrier)
    obj=DispatchBatch.objects.get(pk=batch['id'])
    with django_assert_num_queries(1):
        body=export_batch(obj)
    sheet=load_workbook(BytesIO(body))['SKU']
    assert sheet['B2'].value=='BS-OLD' and sheet['H2'].value==3
    current=load_workbook(BytesIO(export_batch(obj,sku_mode='current')))['SKU']
    assert current['B2'].value==first_line.sku_code_snapshot
    with domain_write():
        second_line.unit_price=None;second_line.save(update_fields=['unit_price'])
    assert load_workbook(BytesIO(export_batch(obj)))['SKU']['H2'].value is None
    assert client(buyer).get(BASE+f'dispatch-batches/{obj.pk}/bigseller/?sku_mode=bad').status_code==400


@pytest.mark.django_db
def test_bigseller_download_scope_and_literal_legacy_code(world):
    from apps.permissions.models import DataScope
    from tests.test_supply_chain_v2_foundation import internal_user
    buyer, first, second, carrier, _=setup(world,mixed_suppliers=True)
    batch,_=create(client(buyer),[first,second],carrier)
    limited=internal_user(world,'bs-limited',scope=DataScope.ScopeType.CUSTOM,config={'supplier_ids':[world['supplier'].pk]})
    assert client(limited).get(BASE+f"dispatch-batches/{batch['id']}/bigseller/").status_code==400
    sku=first.lines.first().sku
    sku.legacy_sku_code='=1+1';sku.save(update_fields=['legacy_sku_code'])
    sheet=load_workbook(BytesIO(export_batch(DispatchBatch.objects.get(pk=batch['id']))))['SKU']
    assert sheet['B2'].value=='=1+1' and sheet['B2'].data_type=='s'


@pytest.mark.django_db
def test_bigseller_inbound_requires_confirmation_and_uses_positive_received_quantity(world):
    buyer, first, _, carrier, warehouse = setup(world)
    api = client(buyer)
    batch, _ = create(api, [first], carrier)
    inbound = generate(api, arrive(api, batch), warehouse)
    endpoint = BASE + f"inbound-orders/{inbound['id']}/bigseller/"
    assert api.get(endpoint).status_code == 400
    inbound, _ = checked(api, inbound, quantities=[0])
    response = write(api, f"inbound-orders/{inbound['id']}/confirm/", {'expected_version': inbound['version']}, 'confirm-bigseller')
    assert response.status_code == 200, response.data
    assert api.get(endpoint).status_code == 400  # all-zero receipt has no exportable rows


@pytest.mark.django_db
def test_bigseller_confirmed_inbound_exports_received_quantity(world):
    buyer, first, _, carrier, warehouse = setup(world)
    api = client(buyer)
    batch, _ = create(api, [first], carrier)
    inbound = generate(api, arrive(api, batch), warehouse)
    inbound, _ = checked(api, inbound, quantities=[2])
    response = write(api, f"inbound-orders/{inbound['id']}/confirm/", {'expected_version': inbound['version']}, 'confirm-bigseller-qty')
    assert response.status_code == 200, response.data
    exported = api.get(BASE + f"inbound-orders/{inbound['id']}/bigseller/")
    assert exported.status_code == 200 and exported.streaming
    sheet = load_workbook(BytesIO(b''.join(exported.streaming_content)))['SKU']
    assert sheet['A2'].value == batch['number'] and sheet['C2'].value == 2


def test_inbound_reconciliation_export_is_blank_counting_template():
    snapshot = {'number': 'IN-1', 'batch_number': 'B-1', 'lines': [
        {'expected_quantity': 5, 'received_quantity': 3, 'product_name': 'Item', 'sku': 'SKU-1'}]}
    sheet = load_workbook(BytesIO(inbound_reconciliation_workbook(snapshot)))['入库清点']
    assert sheet['G2'].value == 5
    assert sheet['H2'].value is None and sheet['I2'].value is None
