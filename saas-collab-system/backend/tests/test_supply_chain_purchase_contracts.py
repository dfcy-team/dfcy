import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient
from apps.tenants.models import Tenant

from apps.supply_chain.models import ShipmentLine, SupplierShipment, domain_write
from tests.test_supply_chain_v2_foundation import confirmed_order, internal_user, world
from tests.test_supply_chain_v2_shipments import INTERNAL, buyer_client, draft, upload, write
from tests.test_supply_chain_v2_numbering import config


@pytest.fixture
def submitted_shipment(world):
    user = internal_user(world, "contract-buyer")
    order = confirmed_order(world, user)
    api = buyer_client(user)
    shipment = draft(api, INTERNAL, world, order, key="contract-shipment")
    shipment = upload(api, INTERNAL + "supplier-shipments/", shipment, key="contract-box")
    response = write(api, INTERNAL + f"supplier-shipments/{shipment['id']}/submit/",
                     {"expected_version": shipment["version"]}, "contract-submit")
    assert response.status_code == 200, response.data
    return user, api, response.data["data"]


def test_contract_draft_idempotent_issue_freezes_source_and_download(submitted_shipment, world):
    user, api, shipment = submitted_shipment
    base = INTERNAL + f"supplier-shipments/{shipment['id']}/purchase-contract/"
    original_line = ShipmentLine.objects.get(shipment_id=shipment["id"], is_active=True)
    before = (original_line.quantity, original_line.unit_price, original_line.product_name_snapshot)
    get_response=api.get(base)
    assert get_response.data["success"] and get_response.data["code"] == "OK"
    defaults = get_response.data["data"]["defaults"]
    defaults["buyer"]["name"] = "买方有限公司"
    defaults["contract_date"] = "2026-10-30"
    defaults["include_order_number"] = False
    defaults["lines"][0]["product_name"] = "=HYPERLINK(\"https://bad\")"
    defaults["lines"][0].update({"unit_price": "25.0000", "tax_rate": "6.000"})
    defaults["pricing_basis"] = "net_plus_tax"
    defaults["tax_treatment"] = "refund"
    body = {"expected_version": 0, "expected_source_version": shipment["version"], "nature": "corporate", "payload": defaults, "reason": "采购合同"}
    saved = api.put(base, body, format="json")
    assert saved.status_code == 200 and saved.data["success"] and "contract" in saved.data["data"], saved.data
    repeated = api.put(base, body, format="json")
    assert repeated.status_code == 200 and repeated.data["data"]["contract"]["id"] == saved.data["data"]["contract"]["id"]
    contract = saved.data["data"]["contract"]
    assert contract["payload"]["lines"][0]["net_total"] == "100.00"
    assert contract["payload"]["lines"][0]["tax_amount"] == "6.00"
    assert contract["payload"]["lines"][0]["line_total"] == "106.00"
    with domain_write():
        source = SupplierShipment.objects.get(pk=shipment["id"]); source.version += 1; source.save()
    stale = api.put(base, body, format="json")
    assert stale.status_code == 409 and stale.data["success"] is False and stale.data["code"] == "CONFLICT"
    assert stale.data["data"]["source_changed"] is True
    shipment["version"] = stale.data["data"]["source_version"]
    body["expected_source_version"] = shipment["version"]
    body["expected_version"] = contract["version"]
    refreshed = api.put(base, body, format="json")
    assert refreshed.status_code == 200 and refreshed.data["data"]["contract"]["source_version"] == shipment["version"]
    contract = refreshed.data["data"]["contract"]
    config(world, user, kind="corporate_contract", prefix="CC")
    rule_kinds=[row["kind"] for row in api.get(INTERNAL + "settings/number-rules/").data["data"]["rules"]]
    assert "corporate_contract" in rule_kinds and "personal_contract" in rule_kinds
    issue_url = base + "issue/"
    issue_body = {"expected_version": contract["version"], "expected_source_version": shipment["version"], "confirmed": True}
    issued = api.post(issue_url, issue_body, format="json")
    assert issued.status_code == 200, issued.data
    assert issued.data["data"]["contract"]["status"] == "issued"
    assert issued.data["data"]["contract"]["number"].startswith("US-CC-")
    repeat_issue = api.post(issue_url, issue_body, format="json")
    assert repeat_issue.status_code == 200 and repeat_issue.data["data"]["contract"]["number"] == issued.data["data"]["contract"]["number"]
    source = ShipmentLine.objects.get(pk=original_line.pk)
    assert (source.quantity, source.unit_price, source.product_name_snapshot) == before
    assert api.put(base, {**body, "expected_version": contract["version"], "payload": {**defaults, "buyer": {"name": "edit"}}}, format="json").status_code == 400
    exported = api.get(base + "download/")
    assert exported.status_code == 200
    workbook = load_workbook(filename=__import__("io").BytesIO(exported.content), data_only=False)
    sheet = workbook.active
    assert sheet["A1"].value == "采 购 合 同"
    assert sheet.max_column == 10 and "B10:D10" in {str(r) for r in sheet.merged_cells.ranges}
    assert sheet["H3"].value.endswith(issued.data["data"]["contract"]["number"])
    assert sheet["A5"].value is None
    assert sheet["A11"].value.startswith("'=")
    assert sheet.max_row >= 11 + 1 + 2 * 10
    assert sheet["G11"].value == "26.5000"
    assert sheet["J11"].value == "106.00" and sheet["H11"].value == "100.00" and sheet["I11"].value == "6.00"
    assert sheet.page_setup.fitToWidth == 1 and sheet.print_area == f"'采购合同'!$A$1:$J${sheet.max_row}"


def test_contract_scope_and_explicit_tax_confirmation(submitted_shipment, world):
    user, api, shipment = submitted_shipment
    base = INTERNAL + f"supplier-shipments/{shipment['id']}/purchase-contract/"
    view = internal_user(world, "contract-viewer", codes=("supply.v2.view",))
    assert buyer_client(view).get(base).status_code == 200
    assert buyer_client(view).put(base, {}, format="json").status_code in {401, 403}
    own = internal_user(world, "contract-own-viewer", scope="own", codes=("supply.v2.view",))
    assert buyer_client(own).get(base).status_code == 404
    other_tenant={"tenant":Tenant.objects.create(name="Other tenant",code="other-contract-tenant")}
    outsider=internal_user(other_tenant,"contract-outsider",codes=("supply.v2.view",))
    assert buyer_client(outsider).get(base).status_code == 404
    data = api.get(base).data["data"]["defaults"]
    data["lines"][0].update({"unit_price": "2.00"})
    result = api.put(base, {"expected_version": 0, "expected_source_version": shipment["version"], "nature": "corporate", "payload": data}, format="json")
    assert result.status_code == 400
    assert not __import__("apps.supply_chain.models", fromlist=["PurchaseContract"]).PurchaseContract.objects.filter(shipment_id=shipment["id"]).exists()
