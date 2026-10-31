from copy import copy, deepcopy
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import hashlib
import json
from datetime import date

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction
from django.http import HttpResponse, FileResponse
from django.utils.http import content_disposition_header
from rest_framework import serializers
from rest_framework.exceptions import NotFound
from rest_framework.response import Response

from .models import (PurchaseContract, ShipmentLine, SupplierShipment, domain_write,
    PurchaseContractSignature, PurchaseContractSignatureAttachment, PurchaseContractSignatureHistory,
    PurchaseContractHistory, contract_lifecycle_write)
from .numbering import number_for
from .permissions import MANAGE, VIEW, internal_orders
from .views import InternalAPI


DEFAULT_TERMS = [
    {"title": "价格保密", "text": "以上合同价格为乙方专供甲方的批量价，为乙方商业机密。甲方不得向双方以外的第三方泄露，否则应承担由此给乙方造成的损失。"},
    {"title": "付款条款和期限", "text": "出货后根据实际发货数量支付货款，并于开具增值税专用发票后三个月内支付货款。"},
    {"title": "交货时间", "text": "乙方应于____年____月____日前交货。"},
    {"title": "交货方法", "text": "请选择：□甲方自提自运；□乙方送货到指定地点。"},
    {"title": "到货地点和接货单位（或接货人）", "text": "甲方如要求变更到货地点或接货人，应在约定交货日前5日书面通知乙方，否则乙方按约定地点交货。"},
    {"title": "双方约定事项", "text": "1. 供方保证按期提供上述产品，不耽误购方采购计划。\n2. 购方收货后24小时内完成检验并核对出库单；发现可能导致产品受损的情况，应于收货后48小时内书面通知供方。"},
    {"title": "货物的技术标准（包括质量要求）", "text": "□按国家标准执行；□无国家标准而有行业标准的，按行业标准执行。"},
    {"title": "其它质量要求", "text": "如有破损、不合格品或商品短缺等异常问题，乙方需免费提供并补发合格商品。"},
    {"title": "包装要求", "text": "1. 纸箱上清晰印刷产品相关内容；\n2. 胶带封箱时箱口封紧，不留空隙；\n3. 保持产品完整，纸箱破损或不符合要求时应重新包装后出货。"},
    {"title": "其他", "text": "本合同一式两份，经双方签字盖章后生效，双方各执一份。收到订单后务必在24小时内回传，否则视为默认。未尽事宜可通过补充协议协商解决。发生经济纠纷，由甲方所在地人民法院管辖。"},
]


class Party(serializers.Serializer):
    name = serializers.CharField(max_length=200, allow_blank=True)
    address = serializers.CharField(max_length=1000, allow_blank=True, required=False, default="")
    representative = serializers.CharField(max_length=200, allow_blank=True, required=False, default="")
    tax_id = serializers.CharField(max_length=100, allow_blank=True, required=False, default="")
    bank = serializers.CharField(max_length=200, allow_blank=True, required=False, default="")
    account = serializers.CharField(max_length=100, allow_blank=True, required=False, default="")
    phone = serializers.CharField(max_length=100, allow_blank=True, required=False, default="")


class ContractLine(serializers.Serializer):
    source_line_id = serializers.IntegerField(min_value=1, allow_null=True, required=False)
    sku = serializers.CharField(max_length=80, allow_blank=True)
    product_name = serializers.CharField(max_length=200, allow_blank=True)
    specification = serializers.CharField(max_length=300, allow_blank=True, required=False, default="")
    unit = serializers.CharField(max_length=40, allow_blank=True, required=False, default="件")
    quantity = serializers.DecimalField(max_digits=14, decimal_places=4, min_value=Decimal("0.0001"))
    unit_price = serializers.DecimalField(max_digits=14, decimal_places=4, min_value=0, required=False, allow_null=True, default=None)
    tax_rate = serializers.DecimalField(max_digits=6, decimal_places=3, min_value=0, max_value=100, required=False, allow_null=True, default=None)
    line_total = serializers.CharField(read_only=True, required=False)
    net_total = serializers.CharField(read_only=True, required=False)
    tax_amount = serializers.CharField(read_only=True, required=False)
    taxed_unit_price = serializers.CharField(read_only=True, required=False)


class ContractTerm(serializers.Serializer):
    title = serializers.CharField(max_length=100, allow_blank=False)
    text = serializers.CharField(max_length=4000, allow_blank=False)


class Payload(serializers.Serializer):
    buyer = Party()
    seller = Party()
    contract_date = serializers.DateField(allow_null=True)
    include_order_number = serializers.BooleanField(default=True)
    tax_treatment = serializers.ChoiceField(choices=["refund", "exempt"], required=False, allow_null=True, default=None)
    pricing_basis = serializers.ChoiceField(choices=["net_plus_tax", "gross_inclusive"], required=False, default="net_plus_tax")
    terms = ContractTerm(many=True, allow_empty=True)
    lines = ContractLine(many=True, allow_empty=False)

    def validate(self, data):
        source_ids = [x.get("source_line_id") for x in data["lines"] if x.get("source_line_id")]
        if len(source_ids) != len(set(source_ids)):
            raise serializers.ValidationError("source_line_id 不可重复。")
        return data


class ContractSave(serializers.Serializer):
    expected_version = serializers.IntegerField(min_value=0)
    expected_source_version = serializers.IntegerField(min_value=1)
    nature = serializers.ChoiceField(choices=["corporate", "personal"])
    payload = Payload()
    reason = serializers.CharField(max_length=500, required=False, allow_blank=True)


def _source(shipment):
    lines = ShipmentLine.objects.filter(shipment=shipment, is_active=True).order_by("line_number")
    return {"buyer": {"name": "", "address": "", "representative": "", "tax_id": "", "bank": "", "account": "", "phone": ""},
        "seller": {"name": shipment.order.supplier.name, "address": "", "representative": "", "tax_id": "", "bank": "", "account": "", "phone": ""},
        "contract_date": None, "include_order_number": True, "pricing_basis": "net_plus_tax", "tax_treatment": None, "terms": [dict(x) for x in DEFAULT_TERMS],
        "lines": [{"source_line_id": x.pk, "sku": x.sku_code_snapshot, "product_name": x.product_name_snapshot,
          "specification": "", "unit": "件", "quantity": str(x.quantity), "unit_price": str(x.unit_price) if x.unit_price is not None else None, "tax_rate": str(getattr(shipment.order.supplier, "default_tax_rate", None)) if getattr(shipment.order.supplier, "default_tax_rate", None) is not None else None} for x in lines]}


def _serialize(contract):
    return {"id": contract.pk, "version": contract.version, "status": contract.status, "number": contract.number,
            "nature": contract.nature, "settlement_tax_amount": str(contract.settlement_tax_amount), "cost_tax_amount": str(contract.cost_tax_amount), "payload": contract.payload, "source_version": contract.source_version}


def _success(data):
    return Response({"data": data})


def _conflict(message, data):
    return Response({"success":False,"code":"CONFLICT","message":message,"data":data}, status=409)


class ContractEndpoint(InternalAPI):
    write_permission_code = MANAGE

    def _shipment(self, request, shipment_id, write=False):
        qs = SupplierShipment.objects.filter(status="submitted", tenant_id=request.user.tenant_id).select_related("order__supplier")
        shipment = qs.filter(pk=shipment_id).first()
        if shipment:
            allowed = internal_orders(request.user, shipment.order.__class__.objects.all(), MANAGE if write else VIEW)
            if not allowed.filter(pk=shipment.order_id).exists(): shipment = None
        if shipment is None:
            raise NotFound("Submitted shipment not found.")
        return shipment

    def get(self, request, shipment_id):
        shipment = self._shipment(request, shipment_id)
        c = PurchaseContract.objects.filter(shipment=shipment).first()
        return _success({"contract": _serialize(c) if c else None, "defaults": _source(shipment), "source_version": shipment.version,
                         "source_changed": bool(c and c.source_version != shipment.version), "shipment_number": shipment.number})

    def put(self, request, shipment_id):
        data = ContractSave(data=request.data); data.is_valid(raise_exception=True)
        shipment = self._shipment(request, shipment_id, write=True)
        try:
            with transaction.atomic():
                shipment = SupplierShipment.objects.select_for_update().get(pk=shipment.pk)
                c = PurchaseContract.objects.select_for_update().filter(shipment=shipment).first()
                expected = data.validated_data["expected_version"]
                source_expected=data.validated_data["expected_source_version"]
                if source_expected != shipment.version:
                    return _conflict("发货单来源已变化，请刷新后确认", {"source_changed":True,"source_version":shipment.version,"contract_source_version":c.source_version if c else None})
                if c and c.status != "draft": raise serializers.ValidationError({"detail":"仅草稿可修改；作废合同请先创建重签草稿。"})
                if (c.version if c else 0) != expected:
                    proposed=Payload().to_representation(data.validated_data["payload"])
                    existing=deepcopy(c.payload) if c else None
                    if existing:
                        for line in existing.get("lines",[]):
                            for key in ("line_total","net_total","tax_amount","taxed_unit_price"): line.pop(key,None)
                    if c and c.status=="draft" and c.nature==data.validated_data["nature"] and existing==proposed:
                        return _success({"contract":_serialize(c),"defaults":_source(shipment),"source_version":shipment.version,"shipment_number":shipment.number})
                    return _conflict("合同草稿版本冲突", {"version": c.version if c else 0})
                p = data.validated_data["payload"]
                if data.validated_data['nature'] == 'corporate' and p.get('pricing_basis') == 'net_plus_tax' and not p.get('tax_treatment'):
                    raise serializers.ValidationError({'payload.tax_treatment': '请选择退税（税费不计成本）或免税（税费计成本）。'})
                if p.get("pricing_basis") == "gross_inclusive" and (not c or c.payload.get("pricing_basis") == "net_plus_tax"):
                    raise serializers.ValidationError({"payload.pricing_basis": "新合同必须使用未税金额加税口径。"})
                raw_basis = request.data.get("payload", {}).get("pricing_basis")
                if c and c.status == "draft" and not c.payload.get("pricing_basis") and raw_basis not in ("gross_inclusive", "net_plus_tax"):
                    raise serializers.ValidationError({"payload.pricing_basis": "此旧草稿使用原含税口径，请明确保持原口径或转换后按未税单价计价。"})
                # Explicit tax-rate confirmation is mandatory; empty means the client omitted a deliberate choice.
                if data.validated_data["nature"] == "personal":
                    for row in p["lines"]: row["tax_rate"] = Decimal("0")
                if any(x["tax_rate"] is None for x in p["lines"]): raise serializers.ValidationError({"payload.lines": "每行必须明确确认税率。"})
                for row in p["lines"]:
                    if row["unit_price"] is None: continue
                    if p.get("pricing_basis") == "net_plus_tax":
                        net = (row["quantity"] * row["unit_price"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                        tax = (net * row["tax_rate"] / 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                        amount = net + tax
                        row["taxed_unit_price"] = str((row["unit_price"] * (1 + row["tax_rate"] / 100)).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP))
                    else:
                        amount = (row["quantity"] * row["unit_price"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                        net = (amount / (1 + row["tax_rate"] / 100)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                        tax = amount - net
                    row["line_total"] = str(amount); row["net_total"] = str(net); row["tax_amount"] = str(tax)
                payload = Payload().to_representation(p)
                for source, rendered in zip(p["lines"], payload["lines"]):
                    for key in ("line_total", "net_total", "tax_amount", "taxed_unit_price"):
                        rendered[key] = source.get(key, "")
                with domain_write():
                    if not c:
                        c = PurchaseContract(tenant=shipment.tenant, created_by=request.user, shipment=shipment,
                            nature=data.validated_data["nature"], payload=payload, source_version=shipment.version)
                    else:
                        c.nature=data.validated_data["nature"]; c.payload=payload; c.source_version=shipment.version; c.version += 1
                    c.last_reason=data.validated_data.get("reason", "")
                    c.save()
                return _success({"contract": _serialize(c), "defaults": _source(shipment), "source_version": shipment.version, "shipment_number": shipment.number})
        except DjangoValidationError as exc: raise serializers.ValidationError(getattr(exc, "message_dict", exc.messages))


class ContractIssue(ContractEndpoint):
    def post(self, request, shipment_id):
        expected = request.data.get("expected_version"); source_expected=request.data.get("expected_source_version")
        if request.data.get("confirmed") is not True: raise serializers.ValidationError({"confirmed": "请确认签发。"})
        shipment=self._shipment(request, shipment_id, write=True)
        with transaction.atomic():
            shipment=SupplierShipment.objects.select_for_update().get(pk=shipment.pk)
            c=PurchaseContract.objects.select_for_update().filter(shipment=shipment).first()
            if not c: raise NotFound("合同草稿不存在，请先保存。")
            if c.status == "issued": return _success({"contract":_serialize(c),"defaults":_source(shipment),"source_version":shipment.version,"shipment_number":shipment.number})
            if c.status != "draft": raise serializers.ValidationError("请先创建重签草稿。")
            if c.payload.get("pricing_basis") not in ("net_plus_tax", "gross_inclusive"):
                raise serializers.ValidationError({"payload.pricing_basis": "请先保存明确选择的合同计价口径。"})
            if not all((c.payload.get("buyer",{}).get("name"), c.payload.get("seller",{}).get("name"), c.payload.get("contract_date"))):
                raise serializers.ValidationError({"detail":"请填写买方、卖方名称和签订日期后再签发。"})
            if shipment.version != source_expected or c.source_version != source_expected: return _conflict("发货单来源已变化", {"source_changed":True,"source_version":shipment.version,"contract_source_version":c.source_version,"shipment_number":shipment.number})
            if c.version != expected: return _conflict("合同草稿版本冲突", {"source_changed":False,"version":c.version})
            if any(x.get("unit_price") is None or x.get("tax_rate") is None for x in c.payload["lines"]):
                raise serializers.ValidationError({"detail":"请先为每项商品填写单价并明确确认税率。"})
            if c.payload.get("pricing_basis") == "net_plus_tax":
                source_lines = {line.pk: line for line in ShipmentLine.objects.filter(shipment=shipment, is_active=True)}
                contract_lines = c.payload.get("lines", [])
                source_ids = [line.get("source_line_id") for line in contract_lines]
                if (len(source_ids) != len(source_lines) or set(source_ids) != set(source_lines)
                        or any(source_id is None for source_id in source_ids)):
                    raise serializers.ValidationError({"payload.lines": "签发前每条发货商品必须且只能对应一条合同商品。"})
                for line in contract_lines:
                    source = source_lines[line["source_line_id"]]
                    if (Decimal(str(line["quantity"])) != source.quantity
                            or Decimal(str(line["unit_price"])) != source.unit_price):
                        raise serializers.ValidationError({"payload.lines": "合同商品数量和未税单价必须与发货单一致。"})
            with domain_write():
                try:
                    prefix="CC" if c.nature=="corporate" else "PC"
                    repeat = c.lifecycle_history.count()
                    fallback = f"{prefix}-{shipment.number}" + (f"-R{repeat+1}" if repeat else "")
                    c.number=number_for(request.user, f"{'corporate' if c.nature=='corporate' else 'personal'}_contract", fallback, country_code=shipment.order.country_code_snapshot)
                except DjangoValidationError as exc:
                    raise serializers.ValidationError(getattr(exc, "message_dict", exc.messages))
                # Corporate tax increases payable; the selected treatment controls goods cost.
                if c.nature == 'corporate' and c.payload.get('pricing_basis') == 'net_plus_tax':
                    if c.payload.get('tax_treatment') not in ('refund', 'exempt'):
                        raise serializers.ValidationError({'payload.tax_treatment': '请先保存退税或免税选择。'})
                    c.settlement_tax_amount = sum((Decimal(x['tax_amount']) for x in c.payload['lines']), Decimal(0))
                    c.cost_tax_amount = c.settlement_tax_amount if c.payload['tax_treatment'] == 'exempt' else Decimal(0)
                c.status="issued"; c.version += 1; c.save()
            return _success({"contract":_serialize(c),"defaults":_source(shipment),"source_version":shipment.version,"shipment_number":shipment.number})


class ContractDownload(ContractEndpoint):
    http_method_names = ["get", "head", "options"]
    def get(self, request, shipment_id, history_id=None):
        shipment=self._shipment(request, shipment_id)
        if history_id is not None:
            from types import SimpleNamespace
            history = PurchaseContractHistory.objects.filter(pk=history_id, tenant_id=request.user.tenant_id, contract__shipment=shipment).first()
            if not history: raise NotFound("历史合同不存在。")
            c = SimpleNamespace(number=history.number, payload=history.snapshot['payload'])
        else:
            c=PurchaseContract.objects.filter(shipment=shipment, status="issued").first()
            if not c: raise NotFound("合同尚未签发。")
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, Side
        from openpyxl.utils import get_column_letter
        wb=Workbook(); ws=wb.active; ws.title="采购合同"
        p=c.payload
        def safe(value):
            return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
        def merged(row, text, *, bold=False, size=10, height=None):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            cell=ws.cell(row,1,safe(text)); cell.font=Font(name="宋体", size=size, bold=bold)
            cell.alignment=Alignment(vertical="center", wrap_text=True)
            if height: ws.row_dimensions[row].height=height
        widths=[22,14,14,14,10,12,15,15,15,17]
        for col,width in enumerate(widths,1): ws.column_dimensions[get_column_letter(col)].width=width
        ws.merge_cells("A1:J1"); ws["A1"]="采 购 合 同"; ws["A1"].font=Font(name="宋体",size=18,bold=True); ws["A1"].alignment=Alignment(horizontal="center")
        ws.row_dimensions[1].height=32
        buyer=p["buyer"]; seller=p["seller"]
        ws["A3"]=safe("甲方（需方）："+buyer.get("name","")); ws["H3"]=safe("合同编号："+c.number)
        ws["A4"]=safe("乙方（供方）："+seller.get("name","")); ws["H4"]=safe("签订时间："+(p.get("contract_date") or ""))
        ws.merge_cells("A3:G3"); ws.merge_cells("H3:J3"); ws.merge_cells("A4:G4"); ws.merge_cells("H4:J4")
        if p.get("include_order_number"): ws["A5"]=safe("订单号："+shipment.number)
        ws.merge_cells("A5:J5")
        merged(7,"甲乙双方在平等自愿、诚实信用的基础上，经双方充分协商，就甲方购买乙方货物事宜达成本合同，以资共同遵守。",height=34)
        merged(9,"一、产品名称，型号，数量，金额等：",bold=True)
        headers=["产品名称","产品规格","","","单位","数量","含税单价","不含税总价","税额","开票含税价"]
        for col,value in enumerate(headers,1): ws.cell(10,col,value)
        ws.merge_cells("B10:D10")
        line_start=11; row_count=max(5,len(p["lines"]))
        thin=Side(style="thin",color="000000")
        for index in range(row_count):
            row=line_start+index; line=p["lines"][index] if index<len(p["lines"]) else {}
            unit_price = line.get("taxed_unit_price") or line.get("unit_price", "")
            values=[line.get("product_name",""),line.get("specification",""),"","",line.get("unit","件"),line.get("quantity",""),unit_price,"","",line.get("line_total","")]
            for col,value in enumerate(values,1): ws.cell(row,col,safe(value))
            ws.merge_cells(start_row=row,start_column=2,end_row=row,end_column=4)
            ws.cell(row,8,safe(line.get("net_total",""))); ws.cell(row,9,safe(line.get("tax_amount","")))
            for col in range(1,11):
                cell=ws.cell(row,col); cell.alignment=Alignment(vertical="center",wrap_text=True); cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
            ws.row_dimensions[row].height=30
        total_row=line_start+row_count
        ws.cell(total_row,1,"合 计 数 量"); ws.cell(total_row,6,safe(str(sum((Decimal(str(x["quantity"])) for x in p["lines"]),Decimal(0)))))
        ws.merge_cells(start_row=total_row,start_column=7,end_row=total_row,end_column=9)
        ws.cell(total_row,7,"合 计 金 额（人民币）："); ws.cell(total_row,10,safe(str(sum((Decimal(str(x.get("line_total",0))) for x in p["lines"]),Decimal(0)))))
        row=total_row+1
        if p.get("pricing_basis") == "net_plus_tax":
            net_sum=sum((Decimal(str(x.get("net_total",0))) for x in p["lines"]),Decimal(0))
            tax_sum=sum((Decimal(str(x.get("tax_amount",0))) for x in p["lines"]),Decimal(0))
            rates=", ".join(sorted({format(Decimal(str(x.get("tax_rate",0))).normalize(), "f") for x in p["lines"]}))
            treatment = "免税（税费计入商品成本）" if p.get("tax_treatment") == "exempt" else "退税（税费不计商品成本）"
            merged(row,f"{treatment}。计价说明：未税业务单价另加税点；适用税率 {rates}%。未税合计 {net_sum:.2f} 元，新增税费 {tax_sum:.2f} 元，含税合计 {net_sum+tax_sum:.2f} 元。")
        row+=2
        for idx,term in enumerate(p.get("terms",[]),1):
            text=f"{idx+1}、{term.get('title','')}: {term.get('text','')}"
            estimated_lines=sum(max(1,(len(part)+89)//90) for part in text.split("\n"))
            merged(row,text,height=max(30,estimated_lines*18+6)); row+=2
        row+=1; ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=5); ws.cell(row,1,"甲方（需方）签章")
        ws.merge_cells(start_row=row,start_column=6,end_row=row,end_column=10); ws.cell(row,6,"乙方（供方）签章")
        for label,key in [("单位名称", "name"),("单位地址", "address"),("法定代表人", "representative"),("纳税人识别号", "tax_id"),("开户银行", "bank"),("账号", "account"),("电话号码", "phone")]:
            row+=1; ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=5); ws.cell(row,1,safe(f"{label}：{buyer.get(key,'')}"))
            ws.merge_cells(start_row=row,start_column=6,end_row=row,end_column=10); ws.cell(row,6,safe(f"{label}：{seller.get(key,'')}"))
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                if cell.value is not None and not cell.font.bold: cell.font=Font(name="宋体",size=10)
                alignment=copy(cell.alignment); alignment.vertical="center"; alignment.wrap_text=True; cell.alignment=alignment
        ws["A1"].alignment=Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.page_setup.orientation = "landscape"
        ws.print_area = f"A1:J{ws.max_row}"
        out=BytesIO(); wb.save(out); out.seek(0)
        resp=HttpResponse(out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"]=f'attachment; filename="{c.number}.xlsx"'; return resp


class ContractHistory(ContractEndpoint):
    http_method_names = ["get", "head", "options"]
    def get(self, request, shipment_id):
        shipment = self._shipment(request, shipment_id)
        rows = PurchaseContractHistory.objects.filter(tenant_id=request.user.tenant_id, contract__shipment=shipment).order_by('-pk')
        try:
            offset = int(request.query_params.get('offset', 0))
            if offset < 0: raise ValueError
        except (ValueError, TypeError): raise serializers.ValidationError('无效分页。')
        result = []
        for row in rows[offset:offset+20]:
            sig = PurchaseContractSignature.objects.filter(pk=row.snapshot.get('signature_id'), tenant_id=request.user.tenant_id).first()
            result.append({'id': row.pk, 'number': row.number, 'version': row.version, 'reason': row.reason,
                'created_at': row.created_at, 'nature': row.snapshot.get('nature'),
                'settlement_tax_amount': row.snapshot.get('settlement_tax_amount'),
                'cost_tax_amount': row.snapshot.get('cost_tax_amount'),
                'signature': _signature_data(sig) if sig else None})
        return _success({'history': result, 'count': rows.count(), 'offset': offset})


class ContractVoid(ContractEndpoint):
    http_method_names = ["post", "options"]
    def post(self, request, shipment_id):
        reason = str(request.data.get('reason') or '').strip()
        if not reason or len(reason) > 500:
            raise serializers.ValidationError({'reason': '作废原因必填，且不超过500字。'})
        shipment = self._shipment(request, shipment_id, write=True)
        from django.db.models import Sum
        from . import costs, services
        from .payment_request_models import SupplyPaymentRequest
        with transaction.atomic():
            services._tenant_lock(request.user)
            shipment = SupplierShipment.objects.select_for_update().get(pk=shipment.pk)
            c = PurchaseContract.objects.select_for_update().filter(shipment=shipment).first()
            if not c or c.status != 'issued':
                raise serializers.ValidationError('仅可作废已签发合同。')
            if type(request.data.get('expected_version')) is not int or c.version != request.data['expected_version']:
                return _conflict('合同版本已变化，请刷新后重试。', {'version': c.version})
            summary = costs.payment_summary(request.user, 'shipment', shipment.pk)['summary']
            reserved = SupplyPaymentRequest.objects.filter(tenant_id=c.tenant_id, kind='shipment', source_id=shipment.pk,
                status='submitted', payment__isnull=True).exclude(approval__status__in=['rejected', 'withdrawn']).exclude(batch__approval__status__in=['rejected', 'withdrawn']).aggregate(v=Sum('amount'))['v'] or Decimal(0)
            if summary['due_amount'] is None:
                raise serializers.ValidationError('当前应付待核对，暂不能作废合同。')
            base = Decimal(summary['due_amount']) - c.settlement_tax_amount
            committed = Decimal(summary['paid_amount']) + Decimal(summary['applied_amount']) + reserved
            if committed > base:
                raise serializers.ValidationError('作废后应付将低于已付款、定金抵扣及有效付款申请，请先冲销或取消相关记录。')
            sig = PurchaseContractSignature.objects.filter(contract=c, contract_number=c.number).first()
            snapshot = {**_serialize(c), 'signature_id': sig.pk if sig else None}
            with domain_write(), contract_lifecycle_write():
                PurchaseContractHistory.objects.create(tenant=c.tenant, created_by=request.user, contract=c,
                    number=c.number, version=c.version, reason=reason, snapshot=snapshot)
                c.status = 'voided'
                c.settlement_tax_amount = 0
                c.cost_tax_amount = 0
                c.version += 1
                c.last_reason = reason
                c.save()
            return _success({'contract': _serialize(c), 'defaults': _source(shipment), 'source_version': shipment.version, 'shipment_number': shipment.number})


class ContractReissue(ContractEndpoint):
    http_method_names = ["post", "options"]
    def post(self, request, shipment_id):
        shipment = self._shipment(request, shipment_id, write=True)
        with transaction.atomic():
            SupplierShipment.objects.select_for_update().get(pk=shipment.pk)
            c = PurchaseContract.objects.select_for_update().filter(shipment=shipment).first()
            if not c or c.status != 'voided':
                raise serializers.ValidationError('仅已作废合同可创建重签草稿。')
            if type(request.data.get('expected_version')) is not int or c.version != request.data['expected_version']:
                return _conflict('合同版本已变化，请刷新后重试。', {'version': c.version})
            with domain_write():
                c.status = 'draft'; c.number = None; c.version += 1; c.save()
            return _success({'contract': _serialize(c), 'defaults': _source(shipment), 'source_version': shipment.version, 'shipment_number': shipment.number})


# Signed copies and their verification state are deliberately separate from the issued contract.
from rest_framework.exceptions import ValidationError
from .payment_vouchers import _validate_file
from .api import key as idempotency_key
from .dates import business_date


def _signature_data(signature):
    attachments = signature.attachments.select_related("created_by").defer("content").order_by("pk") if signature.pk else []
    history = signature.history.select_related("created_by").order_by("-created_at", "-pk") if signature.pk else []
    return {"version": signature.version, "status": signature.status, "signed_date": signature.signed_date,
        "attachments": [{"id": a.pk, "file_name": a.file_name, "media_type": a.media_type, "byte_size": a.byte_size,
            "download_url": f"/api/internal/supply-chain/v2/supplier-shipments/{signature.contract.shipment_id}/purchase-contract/signatures/{a.pk}/download/",
            "uploaded_by": a.created_by.full_name or a.created_by.username, "created_at": a.created_at} for a in attachments],
        "history": [{"status": h.status, "reason": h.reason, "signed_date": h.signed_date, "created_at": h.created_at,
            "uploaded_by": h.created_by.full_name or h.created_by.username, "attachment_ids":h.attachment_ids} for h in history[:50]],
        "history_total": history.count() if signature.pk else 0}


class ContractSignatures(ContractEndpoint):
    def get(self, request, shipment_id):
        shipment = self._shipment(request, shipment_id)
        contract = PurchaseContract.objects.filter(shipment=shipment, status="issued").first()
        if not contract: raise NotFound("合同尚未签发。")
        sig = PurchaseContractSignature.objects.filter(contract=contract, contract_number=contract.number).first()
        if not sig:
            sig = PurchaseContractSignature(contract=contract, tenant=shipment.tenant, created_by=request.user)
        return _success(_signature_data(sig))

    def post(self, request, shipment_id):
        shipment = self._shipment(request, shipment_id, write=True)
        upload = request.FILES.get("file")
        name, media, body, digest = _validate_file(upload)
        contract = PurchaseContract.objects.filter(shipment=shipment, status="issued").first()
        if not contract: raise NotFound("合同尚未签发。")
        original_number = contract.number
        idem = idempotency_key(request)
        with transaction.atomic():
            SupplierShipment.objects.select_for_update().get(pk=shipment.pk)
            contract = PurchaseContract.objects.select_for_update().get(pk=contract.pk)
            if contract.status != "issued" or contract.number != original_number: raise serializers.ValidationError("合同状态已变化，请刷新。")
            sig = PurchaseContractSignature.objects.filter(contract=contract, contract_number=contract.number).first()
            if sig is None:
                with domain_write(): sig = PurchaseContractSignature.objects.create(contract=contract,tenant=shipment.tenant,created_by=request.user)
            old = sig.attachments.filter(idempotency_key=idem).defer("content").first()
            if old:
                if old.sha256 != digest or old.file_name != name:
                    return _conflict("该幂等键已用于不同文件", _signature_data(sig))
                return _success(_signature_data(sig))
            if sig.attachments.count() >= 20: raise serializers.ValidationError("每份合同最多上传20个签署件。")
            with domain_write():
                PurchaseContractSignatureAttachment.objects.create(tenant=shipment.tenant, signature=sig, created_by=request.user,
                    file_name=name, media_type=media, byte_size=len(body), sha256=digest, idempotency_key=idem, content=body)
            sig.version += 1
            if sig.status == "signed":
                sig.status="pending"; sig.signed_date=None
                with domain_write(): PurchaseContractSignatureHistory.objects.create(tenant=shipment.tenant, signature=sig, created_by=request.user,
                    status="pending", signed_date=None, reason="签署件新增，需重新核对", idempotency_key=None,
                    attachment_ids=list(sig.attachments.values_list("pk",flat=True)))
            with domain_write(): sig.save(update_fields=["status","signed_date","version","updated_at"])
            return Response({"data":_signature_data(sig)}, status=201)


class ContractSignatureStatus(ContractEndpoint):
    def post(self, request, shipment_id):
        shipment = self._shipment(request, shipment_id, write=True)
        contract = PurchaseContract.objects.filter(shipment=shipment, status="issued").first()
        if not contract: raise NotFound("合同尚未签发。")
        original_number = contract.number
        status_value=request.data.get("status"); expected=request.data.get("expected_version")
        signed_date=request.data.get("signed_date"); reason=str(request.data.get("reason") or "").strip()
        if status_value not in {"signed", "pending"}: raise serializers.ValidationError({"status":"状态只能为 signed 或 pending。"})
        if not isinstance(expected, int) or isinstance(expected, bool): raise serializers.ValidationError({"expected_version":"必填版本号。"})
        if not reason or len(reason)>500: raise serializers.ValidationError({"reason":"请填写不超过500字的原因。"})
        idem=idempotency_key(request)
        request_hash=hashlib.sha256(json.dumps({"status":status_value,"signed_date":signed_date,"reason":reason,"expected_version":expected},sort_keys=True,ensure_ascii=False).encode()).hexdigest()
        parsed=None
        if status_value=="signed":
            try: parsed=date.fromisoformat(str(signed_date))
            except (ValueError, TypeError): raise serializers.ValidationError({"signed_date":"请提供有效签署日期。"})
            if parsed > business_date(): raise serializers.ValidationError({"signed_date":"签署日期不能晚于今天。"})
        elif signed_date not in (None, ""):
            raise serializers.ValidationError({"signed_date":"撤回签署状态时日期必须为空。"})
        with transaction.atomic():
            SupplierShipment.objects.select_for_update().get(pk=shipment.pk)
            contract=PurchaseContract.objects.select_for_update().get(pk=contract.pk)
            if contract.status != "issued" or contract.number != original_number: raise serializers.ValidationError("合同状态已变化，请刷新。")
            sig=PurchaseContractSignature.objects.select_for_update().filter(contract=contract, contract_number=contract.number).first()
            if sig is None:
                with domain_write(): sig=PurchaseContractSignature.objects.create(contract=contract,tenant=shipment.tenant,created_by=request.user)
            replay=sig.history.filter(idempotency_key=idem).first()
            if replay:
                if replay.request_hash != request_hash: return _conflict("该幂等键已用于其他状态请求",_signature_data(sig))
                return _success(_signature_data(sig))
            if sig.version != expected: return _conflict("签署状态版本冲突",_signature_data(sig))
            if status_value=="signed" and (parsed is None or not sig.attachments.exists()): raise serializers.ValidationError("确认已签署需要签署日期及至少一个签署件。")
            sig.status=status_value; sig.signed_date=parsed if status_value=="signed" else None; sig.version+=1
            with domain_write(): sig.save(update_fields=["status","signed_date","version","updated_at"])
            with domain_write(): PurchaseContractSignatureHistory.objects.create(tenant=shipment.tenant,signature=sig,created_by=request.user,
                status=sig.status,signed_date=sig.signed_date,reason=reason,idempotency_key=idem,request_hash=request_hash,
                attachment_ids=list(sig.attachments.values_list("pk",flat=True)))
            return _success(_signature_data(sig))


class ContractSignatureDownload(ContractEndpoint):
    def get(self, request, shipment_id, attachment_id):
        shipment=self._shipment(request,shipment_id)
        row=PurchaseContractSignatureAttachment.objects.filter(tenant_id=request.user.tenant_id, signature__contract__shipment=shipment,pk=attachment_id).first()
        if row is None: raise NotFound("签署件不存在或无权访问。")
        body=bytes(row.content)
        if hashlib.sha256(body).hexdigest()!=row.sha256: raise serializers.ValidationError("签署件完整性校验失败。")
        response=FileResponse(BytesIO(body),as_attachment=True,filename=row.file_name,content_type=row.media_type)
        response["Content-Disposition"]=content_disposition_header(True,row.file_name)
        response["Cache-Control"]="private, no-store"; response["X-Content-Type-Options"]="nosniff"
        return response
