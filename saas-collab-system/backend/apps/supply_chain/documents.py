"""Deterministic, offline purchase files. Authorization belongs to the API adapter."""
from io import BytesIO
from decimal import Decimal
from datetime import date

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.pdfgen.canvas import Canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.graphics.barcode.qr import QrCodeWidget
from reportlab.graphics.shapes import Drawing
from reportlab.graphics import renderPDF
from reportlab.lib.units import mm


IMPORT_HEADERS = ["导入分组号", "供应商编码", "国家编码", "商品SKU", "采购数量", "采购单价", "备注单号", "下单日期", "备注"]
IMPORT_SHEET = "采购单导入"

PRODUCT_DETAIL_FIELDS = [
    ("spu", "SPU"), ("brand", "品牌"), ("category", "分类"),
    ("specification", "规格"), ("spec_values", "规格值"), ("color", "颜色"),
    ("size", "尺寸"), ("material", "材质"), ("unit", "单位"),
    ("package_length_cm", "包装长cm"), ("package_width_cm", "包装宽cm"),
    ("package_height_cm", "包装高cm"), ("package_weight_g", "包装重量g"),
    ("package_volume_cbm", "包装体积m³"), ("origin_country", "原产国"),
    ("hs_code", "海关编码"), ("inventory_type", "商品形态"),
    ("product_description", "商品描述"), ("selling_points", "卖点"),
    ("product_type", "商品类型"),
]


def _detail_text(value):
    """Render snapshot metadata as literal text, including structured values."""
    if value is None:
        return None
    if isinstance(value, (dict, list, tuple)):
        import json
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _safe_image_link(value):
    """Only expose ordinary web URLs as clickable links in the workbook."""
    from urllib.parse import urlsplit
    try:
        parsed = urlsplit(str(value or ""))
        return parsed.scheme.casefold() in {"http", "https"} and bool(parsed.hostname) and not parsed.username and not parsed.password
    except (TypeError, ValueError):
        return False


def _cell(sheet, row, col, value):
    cell = sheet.cell(row, col, value)
    if isinstance(value, str):
        # Metadata such as reference numbers must never become Excel formulae.
        cell.data_type = "s"
    return cell


def _style(sheet, header_row, widths, *, input_columns=()):
    sheet.sheet_view.showGridLines = False
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(vertical="center", horizontal="right" if cell.data_type == "f" or isinstance(cell.value, (int, float, Decimal)) else "left")
    for cell in sheet[header_row]:
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="28465C")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[header_row].height = 28
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=header_row + 1):
        sheet.row_dimensions[row[0].row].height = 24
        for cell in row:
            if cell.column in input_columns:
                cell.fill = PatternFill("solid", fgColor="FFF3CE")
    sheet.freeze_panes = f"A{header_row + 1}"


def _xlsx(workbook):
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def import_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = IMPORT_SHEET
    sheet.append([f"{label} *" if index < 5 else label for index, label in enumerate(IMPORT_HEADERS)])
    for _ in range(10):
        sheet.append([None] * len(IMPORT_HEADERS))
    _style(sheet, 1, {"A": 18, "B": 20, "C": 18, "D": 26, "E": 16, "F": 16, "G": 24, "H": 18, "I": 38}, input_columns=range(1, 10))
    instructions = workbook.create_sheet("导入说明")
    instructions.sheet_view.showGridLines = False
    instructions.column_dimensions["A"].width = 118
    for text in ["填写说明", "同一分组号生成一张草稿采购单，供应商、国家、备注单号和日期须一致。",
                 "必填：分组号、供应商编码、国家编码、SKU、正整数数量。单价留空会取最近供应商价或商品价。",
                 "备注单号可重复。国家编码使用国家档案中的编码；下单日期格式 YYYY-MM-DD，留空取今天。",
                 "商品SKU支持新编码或旧编码；同一分组内SKU不得重复，跨分组可以重复。带 * 的五列为必填。", "导入只建立草稿，采购核对价格后确认下单。请勿填写公式。"]:
        instructions.append([text])
    instructions["A1"].font = Font(name="Arial", size=15, bold=True)
    positive = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=False)
    positive.errorTitle, positive.error, positive.showErrorMessage = "数量格式", "采购数量须为正整数。", True
    sheet.add_data_validation(positive)
    positive.add("E2:E1001")
    sheet["H2"].number_format = "yyyy-mm-dd"
    return _xlsx(workbook)


def purchase_workbook(snapshot, *, images=None):
    workbook = Workbook()
    purchase = workbook.active
    purchase.title = "采购单"
    supplier_fields = [(key, label) for key, label in PRODUCT_DETAIL_FIELDS if key in {"brand", "specification", "color", "size", "material", "unit", "package_length_cm", "package_width_cm", "package_height_cm", "package_weight_g", "package_volume_cbm"} and any((line.get("product_details") or {}).get(key) not in (None, "", {}, []) for line in snapshot["lines"])]
    metadata = [
        ["供应商采购单", snapshot["number"]],
        ["供应商", snapshot["supplier_name"], "国家", snapshot["country_name"]],
        ["备注单号", snapshot["reference_number"], "下单日期", date.fromisoformat(snapshot["order_date"])],
        ["总数量", snapshot["total_quantity"], "总金额", Decimal(snapshot["total_amount"]) if snapshot["total_amount"] is not None else "价格待完善"],
        ["备注", snapshot["notes"]], [],
        ["商品SKU", "旧SKU", "商品名称", "采购数量", "采购单价", "金额", "商品图片", "图片地址"] + [label for _, label in supplier_fields],
    ]
    for row, values in enumerate(metadata, 1):
        for col, value in enumerate(values, 1):
            _cell(purchase, row, col, value)
    quantities = workbook.create_sheet("发货数量填写")
    boxes = workbook.create_sheet("箱号信息填写")
    for sheet, title in [(quantities, "发货批次数量填写页"), (boxes, "批次发货箱号信息填写页")]:
        _cell(sheet, 1, 1, title)
        _cell(sheet, 2, 1, "采购单号")
        _cell(sheet, 2, 2, snapshot["number"])
    quantities.append(["每批发货复制此页，填写批次号、日期和实际数量；额外 SKU 可续填。"])
    quantities.append(["发货批次号", None, "发货日期", None])
    quantities.append(["商品SKU", "旧SKU", "商品名称", "采购数量", "本批发货数量", "备注"])
    boxes.append(["填写所有箱号后上传此页或自己的装箱文件，系统页面只需上传附件。"])
    boxes.append(["发货批次号", None])
    boxes.append(["箱号", "商品SKU", "箱内数量", "备注"])
    image_rows = []
    for index, line in enumerate(snapshot["lines"], 8):
        for col, value in enumerate([line["sku"], line.get("legacy_sku_code", ""), line["name"], line["quantity"], Decimal(line["unit_price"]) if line["unit_price"] is not None else None], 1):
            _cell(purchase, index, col, value)
        purchase.cell(index, 6, f'=IF(E{index}="","",D{index}*E{index})')
        purchase.cell(index, 5).number_format = "0.0000"
        purchase.cell(index, 6).number_format = "0.00"
        details = line.get("product_details") or {}
        for col, field in enumerate((key for key, _ in supplier_fields), 9):
            value = _detail_text(details.get(field))
            _cell(purchase, index, col, value)
        sku_id = line.get("sku_id")
        image_record = (images or {}).get(sku_id, {})
        body = image_record.get("body") if isinstance(image_record, dict) else None
        status = image_record.get("status", "unavailable") if isinstance(image_record, dict) else "unavailable"
        if body:
            try:
                thumbnail = ExcelImage(BytesIO(body))
                scale = min(85 / thumbnail.width, 75 / thumbnail.height, 1)
                thumbnail.width *= scale
                thumbnail.height *= scale
                purchase.add_image(thumbnail, f"G{index}")
                purchase.row_dimensions[index].height = 85
                image_rows.append(index)
            except (OSError, ValueError, TypeError):
                _cell(purchase, index, 7, "图片读取失败")
        else:
            _cell(purchase, index, 7, "暂无图片" if status == "missing" else "图片读取失败")
        image_url = details.get("image_url")
        if _safe_image_link(image_url):
            link_cell = _cell(purchase, index, 8, "查看原图")
            link_cell.hyperlink = str(image_url)
            link_cell.style = "Hyperlink"
        row = index - 2
        for col, value in enumerate([line["sku"], line.get("legacy_sku_code", ""), line["name"], line["quantity"], None, None], 1):
            _cell(quantities, row, col, value)
    for _ in range(max(10, len(snapshot["lines"]))):
        boxes.append([None] * 4)
    widths = {"A": 28, "B": 22, "C": 48, "D": 18, "E": 22, "F": 18, "G": 20, "H": 20}
    widths.update({get_column_letter(9 + offset): 22 for offset in range(len(supplier_fields))})
    _style(purchase, 7, widths)
    if supplier_fields:
        for row in purchase.iter_rows(min_row=8, min_col=9):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            purchase.row_dimensions[row[0].row].height = 85
    for image_row in image_rows:
        purchase.row_dimensions[image_row].height = 85
    purchase["D3"].number_format = "yyyy-mm-dd"
    _style(quantities, 5, widths, input_columns=(5, 6))
    _style(boxes, 5, {"A": 24, "B": 28, "C": 18, "D": 42}, input_columns=(1, 2, 3, 4))
    for sheet in workbook:
        sheet["A1"].font = Font(name="Arial", size=15, bold=True)
    return _xlsx(workbook)


def inbound_reconciliation_workbook(snapshot, *, images=None):
    """Build one reconciliation row per inbound line, retaining split lines."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "入库清点"
    headers = ["序号", "清点单号", "采购单号", "商品图片", "商品名称", "SKU", "发货数量", "实际清点数量", "差异数量"]
    sheet.append(headers)
    image_rows = []
    for index, line in enumerate(snapshot.get("lines", []), 1):
        expected = line.get("expected_quantity")
        actual = None
        difference = None
        row = sheet.max_row + 1
        for col, value in enumerate([index, snapshot["number"], snapshot["batch_number"], None,
                                     line.get("product_name", ""), line.get("legacy_sku_code") or line.get("sku", ""),
                                     expected, actual, difference], 1):
            _cell(sheet, row, col, value)
        record = (images or {}).get(line.get("sku_id"), {})
        if not isinstance(record, dict):
            record = {}
        body = record.get("body") if isinstance(record, dict) else None
        if body:
            try:
                thumbnail = ExcelImage(BytesIO(body))
                scale = min(80 / thumbnail.width, 70 / thumbnail.height, 1)
                thumbnail.width *= scale
                thumbnail.height *= scale
                sheet.add_image(thumbnail, f"D{row}")
                image_rows.append(row)
            except (OSError, ValueError, TypeError):
                _cell(sheet, row, 4, "图片读取失败")
        else:
            _cell(sheet, row, 4, "暂无图片" if record.get("status") == "missing" else "图片读取失败")
        sheet.cell(row, 9).number_format = '0;[Red]-0;0'
    _style(sheet, 1, {"A": 8, "B": 20, "C": 20, "D": 18, "E": 36, "F": 24, "G": 14, "H": 16, "I": 14})
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 75
    for row in image_rows:
        sheet.row_dimensions[row].height = 75
    return _xlsx(workbook)


def sku_labels(snapshot):
    """One 60x50mm label per distinct SKU. QR and printed code prefer legacy SKU."""
    output = BytesIO()
    canvas = Canvas(output, pagesize=(60 * mm, 50 * mm), pageCompression=1)
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    distinct = {line["sku"]: (str(line.get("legacy_sku_code") or "").strip() or line["sku"], line["name"])
                for line in snapshot["lines"]}
    for sku, (label_code, name) in distinct.items():
        text_lines = []
        for text, font, font_size in [(label_code, "Helvetica-Bold", 8), (name, "STSong-Light", 8)]:
            current = ""
            for char in text:
                if pdfmetrics.stringWidth(current + char, font, font_size) > 54 * mm:
                    text_lines.append((current, font, font_size))
                    current = char
                else:
                    current += char
            if current:
                text_lines.append((current, font, font_size))
        page_height = max(50 * mm, 39 * mm + len(text_lines) * 10)
        canvas.setPageSize((60 * mm, page_height))
        qr = QrCodeWidget(label_code, barLevel="M")
        x0, y0, x1, y1 = qr.getBounds()
        size = 30 * mm
        drawing = Drawing(size, size, transform=[size / (x1 - x0), 0, 0, size / (y1 - y0), 0, 0])
        drawing.add(qr)
        renderPDF.draw(drawing, canvas, 15 * mm, page_height - 33 * mm)
        # Long identifiers/names wrap without truncating or shrinking text.
        y = page_height - 36 * mm
        for text, font, size in text_lines:
            canvas.setFont(font, size)
            canvas.drawCentredString(30 * mm, y, text)
            y -= 10
        canvas.showPage()
    canvas.save()
    return output.getvalue()
