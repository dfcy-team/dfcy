"""Download-only BigSeller purchase sheet, using confirmed inbound quantities."""
from decimal import Decimal, ROUND_HALF_UP
from copy import copy
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError
from .documents import _cell
from .dispatch_allocations import active
from .models import InboundOrder


def export_batch(batch, *, sku_mode="legacy"):
    """Legacy export mapping retained for compatibility; the dispatch API is blocked."""
    if sku_mode not in {"legacy", "current"}:
        raise ValidationError("SKU编码方式须为 legacy 或 current。")
    if batch.status != 'active':
        raise ValidationError('已取消批次不能导出 BigSeller 采购单。')
    allocations = list(active().filter(tenant_id=batch.tenant_id, membership__batch=batch)
        .select_related('shipment_line__sku').order_by('shipment_line_id'))
    grouped = {}
    for allocation in allocations:
        line = allocation.shipment_line
        code = (line.sku.legacy_sku_code or line.sku_code_snapshot) if sku_mode == 'legacy' else line.sku_code_snapshot
        if not code:
            raise ValidationError('存在未维护SKU编码的商品，不能导出。')
        item = grouped.setdefault(code, {'quantity': 0, 'amount': Decimal(0), 'priced': True, 'sku_id': line.sku_id})
        if item['sku_id'] != line.sku_id:
            raise ValidationError(f'{code} 对应多个商品，请改用新SKU编码或先修正旧编码。')
        item['quantity'] += allocation.quantity
        if line.unit_price is None:
            item['priced'] = False
        else:
            item['amount'] += allocation.quantity * line.unit_price
    if not grouped:
        raise ValidationError('批次没有可导出的集货商品。')
    return _workbook(grouped, batch.number, batch.freight_number, sku_mode)


def export_inbound(inbound, *, sku_mode="legacy"):
    if sku_mode not in {"legacy", "current"}:
        raise ValidationError("SKU编码方式须为 legacy 或 current。")
    if inbound.status != 'confirmed':
        raise ValidationError('仅已确认入库单可导出 BigSeller 采购单。')
    rows = list(inbound.lines.select_related('shipment_line__sku').order_by('shipment_line_id'))
    grouped = {}
    for inbound_line in rows:
        if not inbound_line.received_quantity or inbound_line.received_quantity <= 0:
            continue
        line = inbound_line.shipment_line
        # Explicit export mode: legacy matches existing external SKU catalogs; current uses the document snapshot.
        code = (line.sku.legacy_sku_code or line.sku_code_snapshot) if sku_mode == 'legacy' else line.sku_code_snapshot
        if not code:
            raise ValidationError('存在未维护SKU编码的商品，不能导出。')
        item = grouped.setdefault(code, {'quantity': 0, 'amount': Decimal(0), 'priced': True, 'sku_id': line.sku_id})
        if item['sku_id'] != line.sku_id:
            raise ValidationError(f'{code} 对应多个商品，请改用新SKU编码或先修正旧编码。')
        item['quantity'] += inbound_line.received_quantity
        if line.unit_price is None:
            item['priced'] = False
        else:
            item['amount'] += inbound_line.received_quantity * line.unit_price
    if not grouped:
        raise ValidationError('入库单没有正数实收商品可导出。')
    return _workbook(grouped, inbound.batch.number, inbound.batch.freight_number, sku_mode)


def _workbook(grouped, number, freight_number, sku_mode):
    workbook = load_workbook(Path(__file__).parent / 'templates/bigseller-purchase-zh.xlsx')
    sheet = workbook['SKU']
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.value = None
    for index, (code, item) in enumerate(grouped.items(), 2):
        price = (item['amount'] / item['quantity']).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP) if item['priced'] else None
        for column in range(1, 52):
            sheet.cell(index, column)._style = copy(sheet.cell(2, column)._style)
        values = {1: number, 2: code, 3: item['quantity'], 6: 'CNY', 7: 1,
                  8: price, 9: '国内供应商', 11: freight_number or None}
        for column, value in values.items():
            _cell(sheet, index, column, value)
        sheet.cell(index, 8).number_format = '0.0000'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
