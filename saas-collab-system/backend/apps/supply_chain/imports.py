"""Bounded Excel preview/commit. The retained input and SHA identify original rows."""
from datetime import date, datetime
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import re
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from openpyxl.utils.exceptions import InvalidFileException

from django.core.exceptions import ValidationError
from django.db.models import Q
from django.db import transaction
from django.utils import timezone
from .dates import business_date
from openpyxl import load_workbook
from rest_framework.exceptions import NotFound

from apps.masterdata.models import CountrySiteMaster
from apps.products.models import ProductLegacyItem, ProductSKU
from .archives import allowed_suppliers
from .documents import IMPORT_HEADERS, IMPORT_SHEET
from .models import PurchaseImport, SupplierProductPrice, domain_write
from .permissions import MANAGE
from .services import _internal, _tenant_lock, _key, _matching_replay, _version, _event, _prior_event, create_purchase_order

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
SKU_QUERY_BATCH = 400
PRICE_QUERY_BATCH = 300


def _normalized_headers(values):
    """Accept canonical headers with supported trailing required markers."""
    result = []
    for value in values:
        if not isinstance(value, str):
            return None
        header = value.strip()
        previous = None
        while header != previous:
            previous = header
            header = re.sub(r"(?:\s*\*|\s*[（(]必填(?:项)?[）)])+$", "", header).strip()
        result.append(header)
    return result


def _sku_candidates(tenant_id, codes):
    """Fetch current and generated legacy identifiers in bounded queries."""
    current = {}
    legacy = {}
    code_list = list(codes)
    for offset in range(0, len(code_list), SKU_QUERY_BATCH):
        batch = code_list[offset:offset + SKU_QUERY_BATCH]
        rows = ProductSKU.objects.filter(tenant_id=tenant_id).filter(
            Q(sku_code__in=batch) | Q(legacy_sku_code__in=batch)
        ).select_related("spu").only(
            "id", "sku_code", "legacy_sku_code", "is_active", "purchase_price", "spu__product_type", "tenant_id"
        )
        for sku in rows:
            if sku.sku_code in batch:
                current.setdefault(sku.sku_code, {})[sku.pk] = sku
            if sku.legacy_sku_code in batch:
                legacy.setdefault(sku.legacy_sku_code, {})[sku.pk] = sku
        staged = ProductLegacyItem.objects.filter(
            tenant_id=tenant_id, status=ProductLegacyItem.Status.GENERATED,
            legacy_sku_code__in=batch, generated_sku__tenant_id=tenant_id,
        ).select_related("generated_sku__spu").only(
            "legacy_sku_code", "generated_sku_id", "generated_sku__id", "generated_sku__sku_code",
            "generated_sku__legacy_sku_code", "generated_sku__is_active", "generated_sku__purchase_price",
            "generated_sku__spu__product_type",
        )
        for item in staged:
            sku = item.generated_sku
            if item.legacy_sku_code in batch and sku:
                legacy.setdefault(item.legacy_sku_code, {})[sku.pk] = sku
    result = {}
    for code in code_list:
        result[code] = {**current.get(code, {}), **legacy.get(code, {})}
    return result


def _price_candidates(tenant_id, pairs, effective_date):
    """Read latest known prices for only the supplier/SKU pairs in this file."""
    result = {}
    pair_list = list(pairs)
    for offset in range(0, len(pair_list), PRICE_QUERY_BATCH):
        condition = Q(pk__in=[])
        for supplier_id, sku_id in pair_list[offset:offset + PRICE_QUERY_BATCH]:
            condition |= Q(supplier_id=supplier_id, sku_id=sku_id)
        prices = SupplierProductPrice.objects.filter(
            tenant_id=tenant_id, effective_date__lte=effective_date,
        ).filter(condition).order_by("supplier_id", "sku_id", "-effective_date", "-pk")
        for price in prices:
            result.setdefault((price.supplier_id, price.sku_id), price.unit_price)
    return result


def _read_workbook(content):
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValidationError("Excel 文件须在 5MB 以内。")
    try:
        with ZipFile(BytesIO(content)) as archive:
            info = archive.infolist()
            if len(info) > 200 or sum(x.file_size for x in info) > 32 * 1024 * 1024 or any(x.flag_bits & 1 or "vbaproject" in x.filename.lower() for x in info):
                raise ValidationError("文件体积、加密或宏内容不符合导入要求。")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if IMPORT_SHEET not in workbook.sheetnames:
            workbook.close()
            raise ValidationError("请使用模板中的“采购单导入”工作表。")
        sheet = workbook[IMPORT_SHEET]
        if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > len(IMPORT_HEADERS):
            workbook.close()
            raise ValidationError("最多导入 1000 行，列数须与模板一致。")
        return workbook, sheet
    except (BadZipFile, OSError, KeyError, ValueError, ParseError, DefusedXmlException, InvalidFileException) as exc:
        raise ValidationError("无法读取文件，请上传未加密的 .xlsx 模板。") from exc


def preview_excel(actor, content):
    workbook, sheet = _read_workbook(content)
    try:
        header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        if _normalized_headers(header) != IMPORT_HEADERS:
            raise ValidationError("表头必须与下载模板一致。")
        suppliers = {x.code: x for x in allowed_suppliers(actor, MANAGE).filter(status="active")}
        countries = {x.code: x for x in CountrySiteMaster.objects.filter(tenant=actor.tenant, status="active")}
        errors, groups, source_rows, error_count = [], {}, 0, 0
        rows = []
        sku_codes = set()

        for cells in sheet.iter_rows(min_row=2, max_col=len(IMPORT_HEADERS)):
            if all(cell.value is None for cell in cells):
                continue
            values = [cell.value for cell in cells]
            invalid_cells = any(cell.data_type == "f" or cell.data_type == "e" for cell in cells)
            sku_code = str(values[3]).strip() if values[3] is not None and not invalid_cells else ""
            if sku_code:
                sku_codes.add(sku_code)
            rows.append((cells[0].row, values, invalid_cells))
        sku_lookup = _sku_candidates(actor.tenant_id, sku_codes)
        price_pairs = set()
        for _, values, invalid_cells in rows:
            if invalid_cells or len(values) < 6 or values[5] not in (None, ""):
                continue
            supplier = suppliers.get(str(values[1]).strip()) if values[1] is not None else None
            sku_code = str(values[3]).strip() if values[3] is not None else ""
            candidates = sku_lookup.get(sku_code, {})
            if supplier and len(candidates) == 1:
                sku = next(iter(candidates.values()))
                if sku.is_active and sku.spu.product_type == "standard":
                    price_pairs.add((supplier.pk, sku.pk))
        price_lookup = _price_candidates(actor.tenant_id, price_pairs, business_date())

        def error(row, field, message):
            nonlocal error_count
            error_count += 1
            if len(errors) < 20:
                errors.append({"row": row, "field": field, "message": message})

        for row_number, values, invalid_cells in rows:
            source_rows += 1
            if invalid_cells:
                error(row_number, "单元格", "不能使用公式或 Excel 错误值。")
                continue
            field_hint = "数据行"
            try:
                group_key, supplier_code, country_code, sku_code = [str(x).strip() if x is not None else "" for x in values[:4]]
                required = [group_key, supplier_code, country_code, sku_code, values[4]]
                required_fields = IMPORT_HEADERS[:5]
                blank_fields = [required_fields[index] for index, value in enumerate(required)
                                if value is None or (isinstance(value, str) and not value.strip())]
                if blank_fields:
                    for field in blank_fields:
                        error(row_number, field, f"{field}为必填项，请填写后重新上传。")
                    continue
                if len(group_key) > 80:
                    error(row_number, "导入分组号", "导入分组号最多 80 字符。")
                    continue
                supplier = suppliers.get(supplier_code)
                country = countries.get(country_code)
                if not supplier:
                    error(row_number, "供应商编码", "供应商编码不存在、已停用或无权访问，请核对供应商档案。")
                    continue
                if not country:
                    error(row_number, "国家编码", "国家编码不存在或已停用，请核对国家档案。")
                    continue
                candidates = sku_lookup.get(sku_code, {})
                if len(candidates) > 1:
                    error(row_number, "商品SKU", f"SKU 编码“{sku_code}”对应多个商品，请先核对新旧 SKU 映射。")
                    continue
                sku = next(iter(candidates.values()), None)
                if not sku:
                    error(row_number, "商品SKU", f"商品 SKU“{sku_code}”不存在，请核对新 SKU 或旧 SKU 编码。")
                    continue
                if not sku.is_active:
                    error(row_number, "商品SKU", f"商品 SKU“{sku_code}”已停用，请改用有效 SKU。")
                    continue
                if sku.spu.product_type != "standard":
                    error(row_number, "商品SKU", f"商品 SKU“{sku_code}”属于组合商品，组合商品不能参与订货；采购导入仅支持普通商品。")
                    continue
                field_hint = "采购数量"
                qty = Decimal(str(values[4]))
                if isinstance(values[4], bool) or not qty.is_finite() or qty <= 0 or qty != qty.to_integral_value() or qty > 2147483647:
                    raise ValueError("采购数量须为正整数。")
                field_hint = "采购单价"
                price = values[5]
                if price is None or price == "":
                    price = price_lookup.get((supplier.pk, sku.pk))
                    if price is None:
                        price = sku.purchase_price
                if price is not None:
                    price = Decimal(str(price))
                    if not price.is_finite() or price < 0 or price >= Decimal("10000000000") or price != price.quantize(Decimal("0.0001")):
                        raise ValueError("单价须为非负数，最多 10 位整数、4 位小数。")
                field_hint = "下单日期"
                reference, notes = str(values[6] or ""), str(values[8] or "")
                order_date = values[7] or business_date()
                if isinstance(order_date, datetime):
                    order_date = order_date.date()
                if isinstance(order_date, str):
                    order_date = date.fromisoformat(order_date)
                if not isinstance(order_date, date) or order_date > business_date():
                    raise ValueError("下单日期须为 YYYY-MM-DD，且不能晚于今天。")
                field_hint = "备注"
                if len(reference) > 120 or len(notes) > 10000:
                    raise ValueError("备注单号最多 120 字符，备注最多 10000 字符。")
                header_data = {"supplier_id": supplier.pk, "country_id": country.pk, "reference_number": reference,
                               "order_date": str(order_date), "notes": notes}
                group = groups.get(group_key)
                if group and group["payload_header"] != header_data:
                    error(row_number, "导入分组号", "同一分组的供应商、国家、备注单号、日期及备注必须一致。")
                    continue
                if group is None:
                    group = groups[group_key] = {"group_key": group_key, "supplier_name": supplier.name, "country_name": country.name,
                        "reference_number": reference, "payload_header": header_data, "lines": []}
                duplicate = next((x for x in group["lines"] if x["sku_id"] == sku.pk), None)
                if duplicate:
                    error(row_number, "商品SKU", f"同一分组内 SKU“{sku_code}”重复；请检查 Excel 第 {duplicate['import_row_number']} 行和第 {row_number} 行。")
                    continue
                group["lines"].append({"sku_id": sku.pk, "quantity": int(qty), "unit_price": str(price) if price is not None else None, "import_row_number": row_number})
            except (ValueError, TypeError, ArithmeticError) as exc:
                error(row_number, field_hint, str(exc) if isinstance(exc, ValueError) else "数量、单价或日期格式不正确。")
        if not source_rows:
            error(0, "文件", "文件没有采购明细。")
        if len(groups) > 100:
            error(0, "分组", "最多导入 100 张采购单。")
        missing = 0
        total_quantity, total_amount = 0, Decimal(0)
        for group in groups.values():
            lines = group["lines"]
            group["lines_count"] = len(lines)
            group["total_quantity"] = sum(x["quantity"] for x in lines)
            complete = all(x["unit_price"] is not None for x in lines)
            amount = sum((x["quantity"] * Decimal(x["unit_price"]) for x in lines if x["unit_price"] is not None), Decimal(0))
            group["total_amount"] = str(amount) if complete else None
            missing += sum(x["unit_price"] is None for x in lines)
            total_quantity += group["total_quantity"]
            total_amount += amount
        return {"valid": error_count == 0, "rows_count": source_rows, "groups_count": len(groups),
                "error_count": error_count, "errors": errors, "groups": list(groups.values()),
                "missing_prices_count": missing, "total_quantity": total_quantity,
                "total_amount": str(total_amount) if missing == 0 else None}
    except (ParseError, DefusedXmlException, KeyError, IndexError) as exc:
        raise ValidationError("Excel 结构无效，请使用下载的模板重新保存。") from exc
    finally:
        workbook.close()


@transaction.atomic
def create_import(actor, filename, content, *, idempotency_key):
    _internal(actor)
    _tenant_lock(actor)
    key = _key(idempotency_key)
    digest = sha256(content).hexdigest()
    previous = PurchaseImport.objects.filter(tenant=actor.tenant, creation_key=key).first()
    if previous:
        if previous.created_by_id != actor.pk:
            raise NotFound("导入任务不可访问。")
        return _matching_replay(previous, digest)
    if not filename.lower().endswith(".xlsx"):
        raise ValidationError("请上传 .xlsx 文件。")
    preview = preview_excel(actor, content)
    with domain_write():
        return PurchaseImport.objects.create(tenant=actor.tenant, created_by=actor, filename=filename[:255],
            file_sha256=digest, file_content=content, creation_key=key, request_hash=digest, preview=preview)


@transaction.atomic
def commit_import(actor, job_id, *, expected_version, idempotency_key):
    _internal(actor)
    _tenant_lock(actor)
    job = PurchaseImport.objects.filter(tenant=actor.tenant, created_by=actor, pk=job_id).first()
    if job is None:
        raise NotFound("导入任务不可访问。")
    payload = {"expected_version": expected_version}
    if _prior_event(actor, job, "import_committed", idempotency_key, payload):
        return job
    _version(job, expected_version)
    if job.committed_at:
        return job
    if not job.preview["valid"]:
        raise ValidationError("请修正预览错误后重新上传。")
    # Check current authorization/master availability again; prices remain the preview snapshot.
    with domain_write():
        ids = []
        for index, group in enumerate(job.preview["groups"]):
            data = {**group["payload_header"], "source_kind": "import", "import_batch_ref": f"IMPORT-{job.pk}-{job.file_sha256[:12]}", "lines": group["lines"]}
            order = create_purchase_order(actor, data, idempotency_key=f"import-{job.pk}-{index}")
            ids.append(order.pk)
        job.order_ids, job.committed_at, job.version = ids, timezone.now(), job.version + 1
        job.save()
        _event(actor, job, "import_committed", idempotency_key, payload, audit={"order_ids": ids})
    return job


def import_data(job):
    # Only the owner can see a preview or original file. Credentials never enter it.
    preview = {**job.preview, "groups": [{k: v for k, v in group.items() if k not in {"payload_header", "lines"}} for group in job.preview["groups"]]}
    return {"id": job.pk, "version": job.version, "filename": job.filename, "file_sha256": job.file_sha256,
            "committed_at": job.committed_at, "order_ids": job.order_ids, "preview": preview}
